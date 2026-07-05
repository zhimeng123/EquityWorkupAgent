from __future__ import annotations

from pathlib import Path
from io import BytesIO
import re
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt
from docx.text.run import Run
from docx.text.paragraph import Paragraph

from mlc_agent.exceptions import TemplateMappingError
from mlc_agent.schemas import FieldResult


NOT_ASSESSED_TEXT = "Not assessed in MVP1"
VALUE_COLOR = "0070C0"


def _replace_text_in_paragraph(
    paragraph: Paragraph, expected: str, replacement: str
) -> Run:
    full_text = "".join(run.text for run in paragraph.runs)
    start = full_text.find(expected)
    if start < 0:
        raise TemplateMappingError(f"未找到预期模板文本: {expected}")
    end = start + len(expected)
    cursor = 0
    replacement_written = False
    replacement_run: Run | None = None
    for run in paragraph.runs:
        run_start = cursor
        run_end = cursor + len(run.text)
        cursor = run_end
        if run_end <= start or run_start >= end:
            continue
        prefix = run.text[: max(0, start - run_start)] if run_start <= start < run_end else ""
        suffix = run.text[max(0, end - run_start) :] if run_start < end <= run_end else ""
        if not replacement_written:
            run.text = prefix + replacement + suffix
            replacement_written = True
            replacement_run = run
        else:
            run.text = suffix
    if not replacement_written or replacement_run is None:
        raise TemplateMappingError(f"模板文本没有可写入的 run: {expected}")
    return replacement_run


def _set_result_font(run: Run, *, prose: bool = False) -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Arial")
    run.font.all_caps = False
    if prose:
        run.font.size = Pt(10)


def _append_hyperlink(paragraph: Paragraph, url: str) -> None:
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)

    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), VALUE_COLOR)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    run_properties.extend((fonts, color, underline, size))
    run.append(run_properties)
    text = OxmlElement("w:t")
    text.text = url
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _is_blue_run(run: Run) -> bool:
    color = run.font.color.rgb
    return color is not None and str(color).upper() == VALUE_COLOR


def _blue_run_groups(paragraph: Paragraph) -> list[list[Run]]:
    groups: list[list[Run]] = []
    current: list[Run] = []
    for run in paragraph.runs:
        if _is_blue_run(run):
            current.append(run)
        elif current:
            groups.append(current)
            current = []
    if current:
        groups.append(current)
    return [group for group in groups if "".join(run.text for run in group).strip()]


def _replace_blue_group(group: list[Run], replacement: str) -> None:
    original = "".join(run.text for run in group)
    leading = re.match(r"^\s*", original).group(0)
    trailing = re.search(r"\s*$", original).group(0)
    group[0].text = f"{leading}{replacement}{trailing}" if replacement else ""
    _set_result_font(group[0])
    for run in group[1:]:
        run.text = ""


def _mapped_template_elements(
    document: DocumentObject,
    mappings: list[dict[str, Any]],
    field_results: dict[str, FieldResult] | None = None,
) -> tuple[dict[Any, set[str]], dict[Any, set[str]]]:
    mapped_paragraphs: dict[Any, set[str]] = {}
    mapped_cells: dict[Any, set[str]] = {}
    for mapping in mappings:
        for locator in mapping["locators"]:
            expected = locator.get("expected_text")
            result = (field_results or {}).get(mapping["field_id"])
            protected_values = {value for value in (expected, result.value if result else None) if value}
            if field_results is None:
                protected_values.add("*")
            if locator["kind"] == "paragraph":
                element = document.paragraphs[locator["paragraph_index"]]._p
                mapped_paragraphs.setdefault(element, set())
                mapped_paragraphs[element].update(protected_values)
            elif locator["kind"] in {"table_cell", "image_anchor"}:
                cell, _ = _cell_for_locator(document, locator)
                mapped_cells.setdefault(cell._tc, set())
                mapped_cells[cell._tc].update(protected_values)
    return mapped_paragraphs, mapped_cells


def _is_protected_group(group: list[Run], expected_texts: set[str]) -> bool:
    text = "".join(run.text for run in group).strip()
    return "*" in expected_texts or any(
        text in expected or expected in text for expected in expected_texts
    )


def _iter_unique_cells(tables: Any, visited: set[Any] | None = None) -> Any:
    visited = visited if visited is not None else set()
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in visited:
                    continue
                visited.add(cell._tc)
                yield cell, row
                yield from _iter_unique_cells(cell.tables, visited)


def _prevent_row_split(row: Any) -> None:
    row_properties = row._tr.get_or_add_trPr()
    if row_properties.find(qn("w:cantSplit")) is None:
        row_properties.append(OxmlElement("w:cantSplit"))


def _mark_non_mvp_placeholders(
    document: DocumentObject,
    mappings: list[dict[str, Any]],
    field_results: dict[str, FieldResult],
) -> None:
    mapped_paragraphs, mapped_cells = _mapped_template_elements(document, mappings, field_results)

    for paragraph in document.paragraphs:
        expected_texts = mapped_paragraphs.get(paragraph._p, set())
        for group in _blue_run_groups(paragraph):
            if not _is_protected_group(group, expected_texts):
                _replace_blue_group(group, NOT_ASSESSED_TEXT)

    for cell, row in _iter_unique_cells(document.tables):
        expected_texts = mapped_cells.get(cell._tc, set())
        groups = [
            (paragraph, group)
            for paragraph in cell.paragraphs
            for group in _blue_run_groups(paragraph)
            if not _is_protected_group(group, expected_texts)
        ]
        if not groups:
            continue
        has_non_blue_text = any(
            run.text.strip()
            for paragraph in cell.paragraphs
            for run in paragraph.runs
            if not _is_blue_run(run)
        )
        for index, (_, group) in enumerate(groups):
            _replace_blue_group(group, NOT_ASSESSED_TEXT if index == 0 else "")
        if not has_non_blue_text:
            groups[0][0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _prevent_row_split(row)


def find_unmarked_non_mvp_placeholders(
    document: DocumentObject,
    mappings: list[dict[str, Any]],
    field_results: dict[str, FieldResult] | None = None,
) -> list[str]:
    mapped_paragraphs, mapped_cells = _mapped_template_elements(document, mappings, field_results)
    remaining: list[str] = []

    for paragraph in document.paragraphs:
        expected_texts = mapped_paragraphs.get(paragraph._p, set())
        remaining.extend(
            text
            for group in _blue_run_groups(paragraph)
            if not _is_protected_group(group, expected_texts)
            if (text := "".join(run.text for run in group).strip()) != NOT_ASSESSED_TEXT
        )

    for cell, _ in _iter_unique_cells(document.tables):
        expected_texts = mapped_cells.get(cell._tc, set())
        for paragraph in cell.paragraphs:
            remaining.extend(
                text
                for group in _blue_run_groups(paragraph)
                if not _is_protected_group(group, expected_texts)
                if (text := "".join(run.text for run in group).strip())
                != NOT_ASSESSED_TEXT
            )
    return remaining


def _cell_for_locator(document: DocumentObject, locator: dict[str, Any]) -> tuple[Any, Any]:
    table = None
    cell = None
    for depth, step in enumerate(locator["table_path"]):
        tables = document.tables if depth == 0 else cell.tables
        table = tables[step["table_index"]]
        cell = table.cell(step["row_index"], step["column_index"])
    if cell is None or table is None:
        raise TemplateMappingError("表格定位路径不能为空")
    return cell, table


def _paragraph_for_locator(document: DocumentObject, locator: dict[str, Any]) -> Paragraph:
    kind = locator["kind"]
    if kind == "paragraph":
        return document.paragraphs[locator["paragraph_index"]]
    if kind == "header":
        section = document.sections[0]
        header_type = locator.get("header_type", "default")
        if header_type == "default":
            header = section.header
        elif header_type == "first":
            header = section.first_page_header
        elif header_type == "even":
            header = section.even_page_header
        else:
            raise TemplateMappingError(f"不支持的页眉类型: {header_type}")
        return header.paragraphs[locator["paragraph_index"]]
    if kind in {"table_cell", "image_anchor"}:
        cell, _ = _cell_for_locator(document, locator)
        if locator.get("expected_empty"):
            paragraph_index = locator.get("paragraph_index") or 0
            return cell.paragraphs[paragraph_index]
        expected = locator["expected_text"]
        for paragraph in cell.paragraphs:
            if expected in paragraph.text:
                return paragraph
        raise TemplateMappingError(
            f"表格单元格未找到预期文本: table_path={locator['table_path']}"
        )
    raise TemplateMappingError(f"不支持的模板定位类型: {kind}")


def validate_template_mapping(template_path: Path, mappings: list[dict[str, Any]]) -> None:
    document = Document(template_path)
    for mapping in mappings:
        for locator in mapping["locators"]:
            paragraph = _paragraph_for_locator(document, locator)
            if locator.get("expected_empty"):
                if paragraph.text.strip():
                    raise TemplateMappingError(
                        f"字段 {mapping['field_id']} 的空目标校验失败"
                    )
            elif locator["expected_text"] not in paragraph.text:
                raise TemplateMappingError(
                    f"字段 {mapping['field_id']} 的模板原文校验失败: {locator['expected_text']}"
                )


def _write_fixed_table(
    document: DocumentObject, locator: dict[str, Any], rows: Any
) -> None:
    if not isinstance(rows, list) or any(not isinstance(row, list) for row in rows):
        raise TemplateMappingError("fill_fixed_table requires structured_value as list[list]")
    _, table = _cell_for_locator(document, locator)
    start = locator["table_path"][-1]
    for row_offset, values in enumerate(rows):
        target_row = start["row_index"] + row_offset
        if target_row >= len(table.rows):
            raise TemplateMappingError("fixed table data exceeds configured table rows")
        for column_offset, value in enumerate(values):
            target_column = start["column_index"] + column_offset
            if target_column >= len(table.columns):
                raise TemplateMappingError("fixed table data exceeds configured table columns")
            table.cell(target_row, target_column).text = str(value)


def _write_one_locator(
    document: DocumentObject,
    mapping: dict[str, Any],
    locator: dict[str, Any],
    result: FieldResult,
) -> None:
    strategy = mapping["write_strategy"]
    if strategy == "fill_fixed_table":
        _paragraph_for_locator(document, locator)
        _write_fixed_table(document, locator, result.structured_value)
        return
    paragraph = _paragraph_for_locator(document, locator)
    expected = locator.get("expected_text") or ""
    if strategy == "replace_with_hyperlink":
        if not paragraph.text.rstrip().endswith(expected.rstrip()):
            raise TemplateMappingError(f"超链接占位文本必须位于段落末尾: {expected}")
        _replace_text_in_paragraph(paragraph, expected, "")
        _append_hyperlink(paragraph, result.value)
    elif strategy in {"replace_text", "replace_multiline_text"}:
        if locator.get("expected_empty"):
            replacement_run = paragraph.add_run(result.value)
        else:
            replacement_run = _replace_text_in_paragraph(paragraph, expected, result.value)
        if mapping.get("output_style") == "business_description":
            _set_result_font(replacement_run, prose=True)
    elif strategy == "append_after_label":
        if expected not in paragraph.text:
            raise TemplateMappingError(f"未找到保留标签: {expected}")
        _replace_text_in_paragraph(paragraph, expected, expected + result.value)
    elif strategy == "insert_image":
        if not result.artifact_path:
            raise TemplateMappingError("insert_image requires artifact_path")
        image_path = Path(result.artifact_path)
        if not image_path.is_file():
            raise TemplateMappingError(f"图片文件不存在: {image_path}")
        if not locator.get("expected_empty"):
            _replace_text_in_paragraph(paragraph, expected, "")
        paragraph.add_run().add_picture(str(image_path), width=Inches(6.3))
    else:
        raise TemplateMappingError(f"不支持的写入策略: {strategy}")


def write_docx_by_mapping(
    template_path: Path,
    output_path: Path,
    mappings: list[dict[str, Any]],
    field_results: dict[str, FieldResult],
) -> list[tuple[str, str]]:
    document = Document(template_path)
    failures: list[tuple[str, str]] = []
    for mapping in mappings:
        result = field_results.get(mapping["field_id"])
        if result is None:
            continue
        snapshot = BytesIO()
        document.save(snapshot)
        snapshot.seek(0)
        try:
            for locator in mapping["locators"]:
                _write_one_locator(document, mapping, locator, result)
        except (IndexError, KeyError, TemplateMappingError) as exc:
            document = Document(snapshot)
            failures.append((mapping["field_id"], str(exc)))
    _mark_non_mvp_placeholders(document, mappings, field_results)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return failures
