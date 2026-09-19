from __future__ import annotations

import re
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn

from mlc_agent.company_resolver import (
    build_eastmoney_url,
    resolve_a_share_company,
)
from mlc_agent.config import load_source_priorities, load_template_mapping, load_yaml
from mlc_agent.docx_writer import (
    find_unmarked_non_mvp_placeholders,
    validate_template_mapping,
    write_docx_by_mapping,
)
from mlc_agent.eastmoney import (
    fetch_company_profile,
    fetch_financial_period_records,
    fetch_financial_summary,
    fetch_market_snapshot,
    normalize_eastmoney_values,
)
from mlc_agent.exceptions import InputValidationError, WorkupAgentError
from mlc_agent.field_merger import build_static_values, merge_field_values
from mlc_agent.http_client import build_http_client
from mlc_agent.json_io import save_json
from mlc_agent.llm import API_KEY_ENV, get_llm_api_key, summarize_business_description
from mlc_agent.logging_utils import configure_run_logger, get_logger
from mlc_agent.official_site import fetch_official_profile
from mlc_agent.schemas import (
    CompanyIdentity,
    ExecutionStep,
    FailedField,
    FieldResult,
    SelfCheckResult,
    SourceValue,
    WorkupAgentState,
)
from mlc_agent.xueqiu import fetch_xueqiu_snapshot, normalize_xueqiu_values


PLAN_DEFINITIONS = [
    ("validate_template", "Validate the fixed DOCX template and mapping"),
    ("resolve_company", "Resolve one unambiguous A-share company"),
    ("fetch_eastmoney", "Collect Eastmoney profile, financial and quote data"),
    ("fetch_xueqiu", "Collect Xueqiu market data without login state"),
    ("fetch_official_site", "Collect and summarize the official company profile"),
    ("shared_disclosures", "Discover, download and parse company disclosures once"),
    ("part_01", "Collect company supplemental information"),
    ("part_02", "Validate company-specific external links"),
    ("part_03", "Collect United States exposure"),
    ("part_04", "Collect operating performance and standard financial periods"),
    ("part_05", "Select peers and collect related-party transactions"),
    ("part_06", "Calculate liquidity and debt metrics"),
    ("part_07", "Calculate deep financial analysis"),
    ("part_08", "Collect market history and securities analysis"),
    ("part_09", "Generate stock chart artifacts"),
    ("part_10", "Collect governance, shareholders and employees"),
    ("part_11", "Collect audit, changes, litigation and news"),
    ("merge_fields", "Apply configured source priorities and record conflicts"),
    ("write_docx", "Write verified values into the fixed DOCX positions"),
    ("generate_evidence", "Generate evidence, failure, raw data and plan files"),
    ("self_check", "Verify field, evidence and output integrity"),
    ("finalize", "Report final output paths and counts"),
]

EXPECTED_CONFIGURED_FIELD_COUNT = 95
INFORMATIONAL_CHECKS = frozenset({"all_fields_succeeded", "failed_fields_empty"})
_RELATIONSHIP_EMBED_ATTRIBUTE = (
    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
)


def _utc_now() -> datetime:
    return datetime.now(timezone.utc)


def _append_error(state: WorkupAgentState, node: str, message: str) -> list[dict[str, str]]:
    errors = list(state.get("errors", []))
    errors.append({"node": node, "message": message, "captured_at": _utc_now().isoformat()})
    get_logger().warning("%s: %s", node, message)
    return errors


def _candidate_step_status(candidates: list[Any], errors: list[dict[str, Any]]) -> str:
    if errors and candidates:
        return "partial"
    if errors:
        return "failed"
    return "completed"


def _part_field_ids(config_dir: Path, part: str) -> list[str]:
    if not part.startswith("part_"):
        return []
    try:
        data = load_yaml(config_dir / "fields" / f"{part}.yaml")
    except (OSError, ValueError):
        return []
    return [
        item["field_id"]
        for item in data.get("fields", [])
        if isinstance(item, dict) and item.get("field_id")
    ]


_ERROR_NODE_FIELDS = {
    "summarize_business_description": {"business_description"},
}
_ERROR_NODE_SOURCES = {
    "fetch_eastmoney_profile": "eastmoney",
    "fetch_eastmoney_financials": "eastmoney",
    "fetch_eastmoney_periods": "eastmoney",
    "fetch_eastmoney_market": "eastmoney",
    "fetch_xueqiu": "xueqiu",
    "fetch_official_site": "official_site",
}


def _failure_reasons(state: WorkupAgentState) -> dict[str, str]:
    field_ids = [item["field_id"] for item in state.get("field_mapping", [])]
    config_dir = Path(state.get("config_dir", "configs"))
    reasons: dict[str, list[str]] = {field_id: [] for field_id in field_ids}
    try:
        source_priorities = load_source_priorities(config_dir, state.get("field_mapping", []))
    except (OSError, ValueError):
        source_priorities = {}

    def add(field_id: str, message: str) -> None:
        if field_id in reasons and message and message not in reasons[field_id]:
            reasons[field_id].append(message)

    def add_message(node: str, message: str) -> None:
        matched = False
        for field_id in field_ids:
            if message.startswith(f"{field_id}:") or message.startswith(f"{field_id} "):
                add(field_id, f"{node}: {message}")
                matched = True
        for field_id in _ERROR_NODE_FIELDS.get(node, set()):
            add(field_id, f"{node}: {message}")
            matched = True
        source = _ERROR_NODE_SOURCES.get(node)
        if source:
            for field_id, priorities in source_priorities.items():
                if source in priorities:
                    add(field_id, f"{node}: {message}")
                    matched = True
        if not matched:
            for field_id in _part_field_ids(config_dir, node):
                add(field_id, f"{node}: {message}")

    for error in state.get("errors", []):
        if isinstance(error, dict):
            add_message(str(error.get("node") or "unknown"), str(error.get("message") or ""))

    for error in state.get("node_errors", []):
        if isinstance(error, dict):
            add_message(str(error.get("node") or "unknown"), str(error.get("message") or ""))

    for part, raw in state.get("part_results", {}).items():
        if not isinstance(raw, dict):
            continue
        raw_errors = raw.get("errors", [])
        if isinstance(raw_errors, list):
            for error in raw_errors:
                if isinstance(error, dict):
                    field_id = str(error.get("field_id") or "")
                    reason = str(error.get("reason") or error.get("message") or "")
                    if field_id:
                        add(field_id, f"{part}: {reason}")
                    else:
                        add_message(part, reason)
                elif error:
                    add_message(part, str(error))
        if raw.get("status") == "failed" and raw.get("reason"):
            add_message(part, str(raw["reason"]))

    for field_id in field_ids:
        if not reasons[field_id]:
            attempted = ", ".join(source_priorities.get(field_id, []))
            reasons[field_id].append(
                f"No verified candidate returned for {field_id}; "
                f"configured sources attempted: {attempted or 'none'}."
            )
    return {field_id: "; ".join(messages) for field_id, messages in reasons.items()}


def _mark_step(
    state: WorkupAgentState,
    step_id: str,
    status: str,
    detail: str | None = None,
) -> list[dict[str, Any]]:
    steps = [dict(step) for step in state.get("execution_plan", [])]
    for step in steps:
        if step["step_id"] == step_id:
            step["status"] = status
            step["detail"] = detail
            break
    return steps


def initialize_run_node(state: WorkupAgentState) -> dict[str, Any]:
    template_path = Path(state["template_path"]).expanduser().resolve()
    if not template_path.exists() or not template_path.is_file():
        raise InputValidationError(f"模板文件不存在: {template_path}")
    if template_path.suffix.lower() != ".docx":
        raise InputValidationError("MVP1 只支持 .docx 模板。")

    output_root = Path(state["output_root"]).expanduser().resolve()
    company_slug = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff_-]+", "_", state["company_input"]).strip("_")
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    run_id = f"{timestamp}_{company_slug or 'company'}"
    run_dir = output_root / run_id
    run_dir.mkdir(parents=True, exist_ok=False)
    log_path = run_dir / "run.log"
    logger = configure_run_logger(log_path, state.get("debug", False))
    logger.info("Run initialized: %s", run_id)
    return {
        "run_id": run_id,
        "template_path": str(template_path),
        "output_root": str(output_root),
        "run_dir": str(run_dir),
        "created_at": _utc_now().isoformat(),
        "source_values": [],
        "errors": [],
        "field_results": {},
        "evidence_records": [],
        "conflict_records": [],
        "failed_fields": [],
        "step_results": [],
        "report_catalog": [],
        "announcement_catalog": [],
        "part_results": {},
        "artifacts": [],
        "node_errors": [],
        "log_path": str(log_path),
    }


def load_template_mapping_node(state: WorkupAgentState) -> dict[str, Any]:
    config = load_template_mapping(Path(state["config_dir"]))
    mappings = config["fields"]
    validate_template_mapping(Path(state["template_path"]), mappings)
    get_logger().info("Validated %d fixed template mappings", len(mappings))
    return {
        "template_version": config["template_version"],
        "field_mapping": mappings,
        "fillable_fields": [item["field_id"] for item in mappings],
    }


def create_execution_plan_node(state: WorkupAgentState) -> dict[str, Any]:
    plan = [ExecutionStep(step_id=step_id, title=title).model_dump(mode="json") for step_id, title in PLAN_DEFINITIONS]
    plan[0]["status"] = "completed"
    plan[0]["detail"] = f"Validated {len(state['field_mapping'])} fixed positions"
    get_logger().info("Created fixed execution plan with %d steps", len(plan))
    return {"execution_plan": plan, "current_step": "confirm_plan"}


def confirm_plan_node(state: WorkupAgentState) -> dict[str, Any]:
    if state.get("auto_confirm"):
        confirmed = True
    else:
        print("\nExecution plan:")
        for index, step in enumerate(state["execution_plan"], start=1):
            print(f"  {index}. {step['title']}")
        answer = input("Proceed? [y/N]: ").strip().lower()
        confirmed = answer in {"y", "yes"}
    if not confirmed:
        raise WorkupAgentError("用户拒绝执行计划，任务已终止。")
    get_logger().info("Execution plan confirmed")
    return {"confirmed": True}


def resolve_company_node(state: WorkupAgentState) -> dict[str, Any]:
    with build_http_client() as client:
        company = resolve_a_share_company(client, state["company_input"])
    get_logger().info("Resolved company %s (%s)", company.company_short_name, company.stock_code)
    return {
        "company": company.model_dump(mode="json"),
        "current_step": "resolve_company",
        "execution_plan": _mark_step(state, "resolve_company", "completed", company.stock_code),
    }


def fetch_eastmoney_node(state: WorkupAgentState) -> dict[str, Any]:
    company = CompanyIdentity.model_validate(state["company"])
    errors = list(state.get("errors", []))
    initial_error_count = len(errors)
    profile: dict[str, Any] | None = None
    financials: dict[str, Any] | None = None
    period_records: list[dict[str, Any]] = []
    market: dict[str, Any] | None = None
    with build_http_client() as client:
        try:
            profile = fetch_company_profile(client, company)
        except Exception as exc:
            errors = _append_error({**state, "errors": errors}, "fetch_eastmoney_profile", str(exc))
        try:
            financials = fetch_financial_summary(client, company)
        except Exception as exc:
            errors = _append_error({**state, "errors": errors}, "fetch_eastmoney_financials", str(exc))
        try:
            period_records = fetch_financial_period_records(client, company)
        except Exception as exc:
            errors = _append_error({**state, "errors": errors}, "fetch_eastmoney_periods", str(exc))
        try:
            total_shares = None
            if financials:
                total_shares = financials["latest"].get("TOTAL_SHARE")
            market = fetch_market_snapshot(client, company, total_shares)
        except Exception as exc:
            errors = _append_error({**state, "errors": errors}, "fetch_eastmoney_market", str(exc))

    if profile:
        company.company_name = str(profile.get("ORG_NAME") or company.company_name)
        company.company_short_name = str(profile.get("SECURITY_NAME_ABBR") or company.company_short_name)
        company.company_english_name = profile.get("ORG_NAME_EN")
        website = profile.get("ORG_WEB")
        if website:
            website = str(website).strip()
            company.official_website = website if website.startswith(("http://", "https://")) else f"https://{website}"

    captured_at = _utc_now()
    values = normalize_eastmoney_values(company, profile, financials, market, captured_at)
    step_errors = errors[initial_error_count:]
    status = _candidate_step_status(values, step_errors)
    get_logger().info("Eastmoney produced %d verified field candidates", len(values))
    all_values = list(state.get("source_values", [])) + [item.model_dump(mode="json") for item in values]
    return {
        "company": company.model_dump(mode="json"),
        "eastmoney_data": {"profile": profile or {}, "financials": financials or {}, "period_records": period_records, "market": market or {}},
        "source_values": all_values,
        "errors": errors,
        "current_step": "fetch_eastmoney",
        "execution_plan": _mark_step(
            state,
            "fetch_eastmoney",
            status,
            "; ".join(item["message"] for item in step_errors)
            if step_errors
            else f"{len(values)} field candidates",
        ),
    }


def fetch_xueqiu_node(state: WorkupAgentState) -> dict[str, Any]:
    company = CompanyIdentity.model_validate(state["company"])
    errors = list(state.get("errors", []))
    initial_error_count = len(errors)
    quote: dict[str, Any] = {}
    values: list[SourceValue] = []
    try:
        with build_http_client() as client:
            quote = fetch_xueqiu_snapshot(client, company)
        values = normalize_xueqiu_values(company, quote, _utc_now())
        detail = f"{len(values)} field candidates"
    except Exception as exc:
        errors = _append_error({**state, "errors": errors}, "fetch_xueqiu", str(exc))
        detail = "source unavailable"
    step_errors = errors[initial_error_count:]
    status = _candidate_step_status(values, step_errors)
    if step_errors:
        detail = "; ".join(item["message"] for item in step_errors)
    get_logger().info("Xueqiu step completed: %s", detail)
    return {
        "xueqiu_data": quote,
        "source_values": list(state.get("source_values", []))
        + [item.model_dump(mode="json") for item in values],
        "errors": errors,
        "current_step": "fetch_xueqiu",
        "execution_plan": _mark_step(state, "fetch_xueqiu", status, detail),
    }


def fetch_official_site_node(state: WorkupAgentState) -> dict[str, Any]:
    company = CompanyIdentity.model_validate(state["company"])
    profile = state.get("eastmoney_data", {}).get("profile", {})
    errors = list(state.get("errors", []))
    initial_error_count = len(errors)
    values: list[SourceValue] = []
    official_data: dict[str, Any] = {}
    description_source = "eastmoney"
    source_url = build_eastmoney_url(company)
    source_text = "\n".join(
        str(value).strip()
        for value in (profile.get("ORG_PROFILE"), profile.get("BUSINESS_SCOPE"))
        if value
    )

    if company.official_website:
        try:
            with build_http_client() as client:
                page = fetch_official_profile(client, company.official_website)
            official_data = {"profile_url": page.url, "profile_text": page.text}
            description_source = "official_site"
            source_url = page.url
            source_text = page.text
            values.append(
                SourceValue(
                    field_id="official_website",
                    value=company.official_website,
                    raw_value=company.official_website,
                    source="official_site",
                    source_url=page.url,
                    captured_at=_utc_now(),
                )
            )
        except Exception as exc:
            errors = _append_error({**state, "errors": errors}, "fetch_official_site", str(exc))

    if source_text:
        try:
            api_key = get_llm_api_key()
        except ValueError:
            errors = _append_error(
                {**state, "errors": errors}, "summarize_business_description", f"{API_KEY_ENV} is not set"
            )
        else:
            try:
                description = summarize_business_description(
                    api_key,
                    source_text,
                    company_name=company.company_english_name or company.company_name,
                    founded_date=profile.get("FOUND_DATE"),
                    listing_date=profile.get("LISTING_DATE"),
                    main_business=profile.get("MAIN_BUSINESS"),
                )
                supporting_sources = []
                if description_source == "official_site":
                    supporting_sources.append(
                        {"source": "eastmoney", "source_url": build_eastmoney_url(company)}
                    )
                values.append(
                    SourceValue(
                        field_id="business_description",
                        value=description,
                        raw_value=source_text,
                        source=description_source,
                        source_url=source_url,
                        captured_at=_utc_now(),
                        metadata={"supporting_sources": supporting_sources},
                    )
                )
            except Exception as exc:
                errors = _append_error({**state, "errors": errors}, "summarize_business_description", str(exc))
    else:
        errors = _append_error(
            {**state, "errors": errors}, "fetch_official_site", "No official or Eastmoney profile text available"
        )

    get_logger().info("Official-site step produced %d verified field candidates", len(values))
    step_errors = errors[initial_error_count:]
    status = _candidate_step_status(values, step_errors)

    return {
        "official_site_data": official_data,
        "source_values": list(state.get("source_values", []))
        + [item.model_dump(mode="json") for item in values],
        "errors": errors,
        "current_step": "fetch_official_site",
        "execution_plan": _mark_step(
            state,
            "fetch_official_site",
            status,
            "; ".join(item["message"] for item in step_errors)
            if step_errors
            else f"{len(values)} field candidates",
        ),
    }


def merge_fields_node(state: WorkupAgentState) -> dict[str, Any]:
    company = CompanyIdentity.model_validate(state["company"])
    candidates = [SourceValue.model_validate(item) for item in state.get("source_values", [])]
    candidates.extend(build_static_values(company, Path(state["template_path"]), _utc_now()))
    priorities = load_source_priorities(Path(state["config_dir"]), state["field_mapping"])
    results, evidence, conflicts, failures = merge_field_values(
        state["field_mapping"], priorities, candidates, _failure_reasons(state)
    )
    get_logger().info(
        "Merged fields: %d selected, %d failed, %d conflicts",
        len(results),
        len(failures),
        len(conflicts),
    )
    return {
        "field_results": {key: value.model_dump(mode="json") for key, value in results.items()},
        "evidence_records": [item.model_dump(mode="json") for item in evidence],
        "conflict_records": [item.model_dump(mode="json") for item in conflicts],
        "failed_fields": [item.model_dump(mode="json") for item in failures],
        "source_values": [item.model_dump(mode="json") for item in candidates],
        "current_step": "merge_fields",
        "execution_plan": _mark_step(state, "merge_fields", "completed", f"{len(results)} selected fields"),
    }


def write_docx_node(state: WorkupAgentState) -> dict[str, Any]:
    output_path = Path(state["run_dir"]) / "result.docx"
    results = {key: FieldResult.model_validate(value) for key, value in state["field_results"].items()}
    write_failures = write_docx_by_mapping(
        Path(state["template_path"]), output_path, state["field_mapping"], results
    )
    get_logger().info("DOCX written to %s with %d write failures", output_path, len(write_failures))
    failed_ids = {field_id for field_id, _ in write_failures}
    if failed_ids:
        mappings = {item["field_id"]: item for item in state["field_mapping"]}
        failed_fields = list(state["failed_fields"])
        for field_id, reason in write_failures:
            failed_fields.append(
                FailedField(
                    field_id=field_id,
                    label=mappings[field_id]["label"],
                    reason=f"DOCX write failed: {reason}",
                    source_attempted=[state["field_results"][field_id]["selected_source"]],
                ).model_dump(mode="json")
            )
        field_results = {
            key: value for key, value in state["field_results"].items() if key not in failed_ids
        }
        evidence = [item for item in state["evidence_records"] if item["field_id"] not in failed_ids]
    else:
        failed_fields = state["failed_fields"]
        field_results = state["field_results"]
        evidence = state["evidence_records"]
    return {
        "output_docx_path": str(output_path),
        "field_results": field_results,
        "evidence_records": evidence,
        "failed_fields": failed_fields,
        "current_step": "write_docx",
        "execution_plan": _mark_step(state, "write_docx", "completed", str(output_path)),
    }


def generate_evidence_files_node(state: WorkupAgentState) -> dict[str, Any]:
    run_dir = Path(state["run_dir"])
    sources_path = run_dir / "sources.json"
    failed_path = run_dir / "failed_fields.json"
    extracted_path = run_dir / "extracted_data.json"
    plan_path = run_dir / "execution_plan.json"
    save_json(sources_path, {"evidence": state["evidence_records"], "conflicts": state["conflict_records"]})
    save_json(failed_path, state["failed_fields"])
    save_json(
        extracted_path,
        {
            "run": {
                "run_id": state["run_id"],
                "goal": state["goal"],
                "created_at": state["created_at"],
                "template_path": state["template_path"],
                "template_version": state["template_version"],
            },
            "company": state["company"],
            "eastmoney": state.get("eastmoney_data", {}),
            "xueqiu": state.get("xueqiu_data", {}),
            "official_site": state.get("official_site_data", {}),
            "part_results": state.get("part_results", {}),
            "artifacts": state.get("artifacts", []),
            "node_errors": [*state.get("errors", []), *state.get("node_errors", [])],
        },
    )
    plan = _mark_step(state, "generate_evidence", "completed", "JSON artifacts written")
    save_json(plan_path, plan)
    get_logger().info("Evidence artifacts written to %s", run_dir)
    return {
        "sources_json_path": str(sources_path),
        "failed_fields_json_path": str(failed_path),
        "extracted_data_json_path": str(extracted_path),
        "execution_plan_json_path": str(plan_path),
        "execution_plan": plan,
        "current_step": "generate_evidence",
    }


def _word_xml_text(element: Any) -> str:
    """Reconstruct visible Word text, including explicit line-break nodes."""
    text: list[str] = []
    for node in element.iter():
        if node.tag == qn("w:t"):
            text.append(node.text or "")
        elif node.tag in {qn("w:br"), qn("w:cr")}:
            text.append("\n")
    return "".join(text)


def _value_is_written(document: Any, mapping: dict[str, Any], value: str) -> bool:
    if mapping["write_strategy"] != "insert_image" and not value.strip():
        return False

    def paragraph_text(paragraph: Any) -> str:
        return _word_xml_text(paragraph._p)

    def cell_for(locator: dict[str, Any]) -> Any:
        cell = None
        for depth, step in enumerate(locator["table_path"]):
            tables = document.tables if depth == 0 else cell.tables
            cell = tables[step["table_index"]].cell(step["row_index"], step["column_index"])
        return cell

    def paragraph_has_embedded_image(paragraph: Any) -> bool:
        for blip in paragraph._p.xpath(".//w:drawing//a:blip"):
            relationship_id = blip.get(_RELATIONSHIP_EMBED_ATTRIBUTE)
            if (
                relationship_id in paragraph.part.rels
                and paragraph.part.rels[relationship_id].reltype.endswith("/image")
            ):
                return True
        return False

    for locator in mapping["locators"]:
        kind = locator["kind"]
        if kind == "paragraph":
            paragraph = document.paragraphs[locator["paragraph_index"]]
            found = (
                paragraph_has_embedded_image(paragraph)
                if mapping["write_strategy"] == "insert_image"
                else value in paragraph_text(paragraph)
            )
        elif kind == "header":
            section = document.sections[0]
            header = {
                "default": section.header,
                "first": section.first_page_header,
                "even": section.even_page_header,
            }[locator.get("header_type", "default")]
            paragraph = header.paragraphs[locator["paragraph_index"]]
            found = (
                paragraph_has_embedded_image(paragraph)
                if mapping["write_strategy"] == "insert_image"
                else value in paragraph_text(paragraph)
            )
        elif kind in {"table_cell", "image_anchor"}:
            cell = cell_for(locator)
            found = (
                any(paragraph_has_embedded_image(paragraph) for paragraph in cell.paragraphs)
                if mapping["write_strategy"] == "insert_image"
                else any(value in paragraph_text(paragraph) for paragraph in cell.paragraphs)
            )
        else:
            found = False
        if not found:
            return False
    return True


def _fixed_table_is_written(
    document: Any,
    mapping: dict[str, Any],
    structured_value: Any,
) -> bool:
    if not isinstance(structured_value, list) or not structured_value:
        return False
    if any(
        not isinstance(row, list)
        or not row
        or any(value is None or not str(value).strip() for value in row)
        for row in structured_value
    ):
        return False
    for locator in mapping["locators"]:
        table = None
        cell = None
        for depth, step in enumerate(locator["table_path"]):
            tables = document.tables if depth == 0 else cell.tables
            table = tables[step["table_index"]]
            cell = table.cell(step["row_index"], step["column_index"])
        if table is None:
            return False
        start = locator["table_path"][-1]
        for row_offset, values in enumerate(structured_value):
            for column_offset, value in enumerate(values):
                row_index = start["row_index"] + row_offset
                column_index = start["column_index"] + column_offset
                if row_index >= len(table.rows) or column_index >= len(table.columns):
                    return False
                if table.cell(row_index, column_index).text != str(value):
                    return False
    return True


def self_check_node(state: WorkupAgentState) -> dict[str, Any]:
    output_path = Path(state["output_docx_path"])
    issues: list[str] = []
    checks: dict[str, bool] = {}
    checks["result_docx_exists"] = output_path.exists() and output_path.stat().st_size > 0
    configured_ids = {item["field_id"] for item in state["field_mapping"]}
    success_ids = set(state["field_results"])
    failure_ids = {item["field_id"] for item in state["failed_fields"]}
    checks["configured_field_count_is_95"] = (
        len(state["field_mapping"]) == EXPECTED_CONFIGURED_FIELD_COUNT
        and len(configured_ids) == EXPECTED_CONFIGURED_FIELD_COUNT
    )
    checks["all_fields_succeeded"] = (
        len(state["field_results"]) == EXPECTED_CONFIGURED_FIELD_COUNT
        and success_ids == configured_ids
    )
    checks["failed_fields_empty"] = not state["failed_fields"]
    checks["conflict_records_empty"] = not state.get("conflict_records", [])
    checks["every_configured_field_accounted_for"] = configured_ids == success_ids | failure_ids
    checks["success_and_failure_disjoint"] = not success_ids & failure_ids
    evidence_by_id = {
        item.get("field_id"): item
        for item in state["evidence_records"]
        if isinstance(item, dict) and item.get("field_id")
    }
    evidence_ids = set(evidence_by_id)
    checks["evidence_matches_field_results"] = (
        len(evidence_by_id) == len(state["evidence_records"])
        and evidence_ids == success_ids
        and all(
            evidence_by_id[field_id].get("value") == result.get("value")
            and evidence_by_id[field_id].get("normalized_value") == result.get("value")
            for field_id, result in state["field_results"].items()
        )
    )
    checks["source_values_configured"] = all(
        isinstance(item, dict) and item.get("field_id") in configured_ids
        for item in state.get("source_values", [])
    )
    checks["every_success_has_evidence"] = success_ids == evidence_ids
    checks["every_failure_has_reason"] = all(item.get("reason") for item in state["failed_fields"])
    checks["every_evidence_has_source_url"] = all(
        item.get("source_url") for item in state["evidence_records"]
    )
    checks["json_outputs_exist"] = all(
        Path(state[key]).exists()
        for key in (
            "sources_json_path",
            "failed_fields_json_path",
            "extracted_data_json_path",
            "execution_plan_json_path",
        )
    )
    checks["json_outputs_parse"] = checks["json_outputs_exist"]
    if checks["json_outputs_exist"]:
        try:
            for key in (
                "sources_json_path",
                "failed_fields_json_path",
                "extracted_data_json_path",
                "execution_plan_json_path",
            ):
                with Path(state[key]).open("r", encoding="utf-8") as handle:
                    json.load(handle)
        except (OSError, json.JSONDecodeError):
            checks["json_outputs_parse"] = False
    checks["artifacts_exist"] = all(
        Path(item["path"]).is_file() for item in state.get("artifacts", [])
    )
    artifact_fields = {
        item.get("field_id") for item in state.get("artifacts", []) if item.get("field_id")
    }
    successful_artifact_fields = {
        field_id
        for field_id, result in state.get("field_results", {}).items()
        if result.get("artifact_path")
    }
    checks["artifact_fields_have_results"] = (
        artifact_fields - {"negative_news_attachment"}
    ) <= successful_artifact_fields
    checks["result_artifacts_registered"] = all(
        any(Path(item["path"]).resolve() == Path(result["artifact_path"]).resolve() for item in state.get("artifacts", []))
        for result in state.get("field_results", {}).values()
        if result.get("artifact_path")
    )

    checks["written_values_match_mapping"] = False
    document = None
    if checks["result_docx_exists"]:
        try:
            document = Document(output_path)
        except Exception as exc:
            issues.append(f"DOCX cannot be reopened: {exc}")
    if document is not None:
        mapping_by_id = {item["field_id"]: item for item in state["field_mapping"]}
        mismatches = []
        for field_id, result in state["field_results"].items():
            mapping = mapping_by_id.get(field_id)
            if mapping is None:
                mismatches.append(field_id)
                continue
            if mapping["write_strategy"] == "fill_fixed_table":
                written = _fixed_table_is_written(
                    document,
                    mapping,
                    result.get("structured_value"),
                )
            else:
                written = _value_is_written(document, mapping, result["value"])
            if not written:
                mismatches.append(field_id)
        checks["written_values_match_mapping"] = not mismatches
        if mismatches:
            issues.append(f"DOCX value mismatch: {', '.join(mismatches)}")
        failed_mapping = {
            item["field_id"]: item for item in state["field_mapping"] if item["field_id"] in failure_ids
        }
        failed_targets_preserved = True
        for mapping in failed_mapping.values():
            for locator in mapping["locators"]:
                expected = locator.get("expected_text")
                if expected and not _value_is_written(
                    document,
                    {**mapping, "locators": [locator], "write_strategy": "replace_text"},
                    expected,
                ):
                    failed_targets_preserved = False
                if locator.get("expected_empty"):
                    kind = locator["kind"]
                    if kind == "paragraph":
                        target = document.paragraphs[locator["paragraph_index"]]
                        is_empty = not target.text.strip() and not target._p.xpath(".//a:blip")
                    elif kind == "header":
                        section = document.sections[0]
                        header = {
                            "default": section.header,
                            "first": section.first_page_header,
                            "even": section.even_page_header,
                        }[locator.get("header_type", "default")]
                        target = header.paragraphs[locator["paragraph_index"]]
                        is_empty = not target.text.strip() and not target._p.xpath(".//a:blip")
                    else:
                        cell = None
                        for depth, step in enumerate(locator["table_path"]):
                            tables = document.tables if depth == 0 else cell.tables
                            cell = tables[step["table_index"]].cell(
                                step["row_index"], step["column_index"]
                            )
                        is_empty = not cell.text.strip() and not cell._tc.xpath(".//a:blip")
                    if not is_empty:
                        failed_targets_preserved = False
        checks["failed_targets_preserved"] = failed_targets_preserved
        image_mappings = [
            mapping
            for mapping in state["field_mapping"]
            if mapping["write_strategy"] == "insert_image"
            and mapping["field_id"] in success_ids
        ]
        checks["images_embedded"] = all(
            _value_is_written(document, mapping, state["field_results"][mapping["field_id"]]["value"])
            for mapping in image_mappings
        )
        remaining_placeholders = find_unmarked_non_mvp_placeholders(
            document,
            state["field_mapping"],
            {
                key: FieldResult.model_validate(value)
                for key, value in state["field_results"].items()
            },
        )
        checks["non_mvp_placeholders_marked"] = not remaining_placeholders
        if remaining_placeholders:
            issues.append(
                "Unmarked non-MVP placeholders: " + "; ".join(remaining_placeholders[:5])
            )
        checks["docx_reopens"] = True
    else:
        checks["failed_targets_preserved"] = False
        checks["images_embedded"] = False
        checks["docx_reopens"] = False

    for name, passed in checks.items():
        if name in INFORMATIONAL_CHECKS:
            continue
        if not passed and not any(name in issue for issue in issues):
            issues.append(f"Self-check failed: {name}")
    result = SelfCheckResult(
        passed=all(
            passed for name, passed in checks.items() if name not in INFORMATIONAL_CHECKS
        ),
        checks=checks,
        issues=issues,
    )
    plan = _mark_step(state, "self_check", "completed" if result.passed else "failed", "; ".join(issues))
    save_json(Path(state["execution_plan_json_path"]), plan)
    get_logger().info("Self-check passed=%s", result.passed)
    return {
        "self_check_result": result.model_dump(mode="json"),
        "execution_plan": plan,
        "current_step": "self_check",
    }


def finalize_run_node(state: WorkupAgentState) -> dict[str, Any]:
    success_count = len(state["field_results"])
    failed_count = len(state["failed_fields"])
    self_check_passed = bool(state.get("self_check_result", {}).get("passed"))
    get_logger().info("Run finalized: %d success, %d failed", success_count, failed_count)
    print(f"\nRun directory: {state['run_dir']}")
    print(f"Successful fields: {success_count}")
    print(f"Failed fields: {failed_count}")
    print(f"Result DOCX: {state['output_docx_path']}")
    plan = _mark_step(
        state,
        "finalize",
        "completed" if self_check_passed else "failed",
        "Run completed" if self_check_passed else "Run blocked by failed self-check",
    )
    save_json(Path(state["execution_plan_json_path"]), plan)
    return {
        "current_step": "finalize",
        "execution_plan": plan,
    }
