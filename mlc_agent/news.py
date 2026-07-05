from __future__ import annotations

import re
from datetime import date
from pathlib import Path
from typing import Literal
from urllib.parse import urlparse

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pydantic import BaseModel, Field, HttpUrl, model_validator

from mlc_agent.litigation import subtract_years


SEARCH_ENGINE_DOMAINS = (
    "google.com",
    "bing.com",
    "baidu.com",
    "duckduckgo.com",
    "sogou.com",
    "so.com",
)


class NewsArticle(BaseModel):
    company_name: str = Field(min_length=1)
    published_date: date
    publisher: str = Field(min_length=1)
    language: Literal["zh", "en"]
    title: str = Field(min_length=1)
    original_url: HttpUrl
    source_kind: Literal["media", "company", "exchange", "regulator"]
    title_entities: list[str] = Field(min_length=1)
    fact_fingerprint: str = Field(min_length=1)
    company_specific_adverse_fact: Literal[True]
    body: str | None = None
    english_summary: str | None = None

    @model_validator(mode="after")
    def validate_article(self) -> "NewsArticle":
        host = (urlparse(str(self.original_url)).hostname or "").lower()
        if any(host == domain or host.endswith("." + domain) for domain in SEARCH_ENGINE_DOMAINS):
            raise ValueError("search-engine URLs are discovery-only and cannot be evidence")
        if self.body is None and self.english_summary is not None:
            raise ValueError("an article without body text cannot have a generated summary")
        if self.body is not None and not self.body.strip():
            raise ValueError("article body cannot be blank")
        if self.body is not None and not self.english_summary:
            raise ValueError("an article with body text requires a verified English summary")
        return self


class NewsEvent(BaseModel):
    event_key: str
    articles: list[NewsArticle]


def _normalize(value: str) -> str:
    return re.sub(r"\W+", "", value, flags=re.UNICODE).casefold()


def news_event_key(article: NewsArticle) -> str:
    entities = ",".join(sorted(_normalize(item) for item in article.title_entities))
    return "|".join(
        (_normalize(article.company_name), article.published_date.isoformat(), entities, _normalize(article.fact_fingerprint))
    )


def select_and_deduplicate_news(articles: list[NewsArticle], *, as_of: date) -> list[NewsEvent]:
    window_start = subtract_years(as_of, 1)
    grouped: dict[str, list[NewsArticle]] = {}
    seen_urls: set[str] = set()
    for article in articles:
        if not window_start <= article.published_date <= as_of:
            continue
        url = str(article.original_url)
        if url in seen_urls:
            continue
        seen_urls.add(url)
        grouped.setdefault(news_event_key(article), []).append(article)
    return [
        NewsEvent(
            event_key=key,
            articles=sorted(items, key=lambda item: (item.body is None, str(item.original_url))),
        )
        for key, items in sorted(grouped.items(), key=lambda pair: pair[1][0].published_date, reverse=True)
    ]


def _set_font(run, *, size: float, bold: bool = False, color: str = "000000") -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((color, underline))
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.extend((properties, text_element))
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def build_negative_news_attachment(
    events: list[NewsEvent],
    *,
    output_path: Path,
    company_name: str,
    as_of: date,
) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_font(header.add_run("Negative News Source Packet"), size=9, color="666666")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_font(footer.add_run(f"Prepared as of {as_of.isoformat()}"), size=9, color="666666")

    kicker = document.add_paragraph()
    kicker.paragraph_format.space_after = Pt(8)
    _set_font(kicker.add_run("SOURCE PACKET"), size=10, bold=True, color="7A5A00")
    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    _set_font(title.add_run("Negative News Articles"), size=24, bold=True)
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    _set_font(subtitle.add_run(f"{company_name} | 12-month review through {as_of.isoformat()}"), size=12, color="555555")

    for event_number, event in enumerate(events, start=1):
        event_heading = document.add_paragraph(style="Heading 1")
        event_heading.add_run(f"Event {event_number}")
        for source_number, article in enumerate(event.articles, start=1):
            source_heading = document.add_paragraph(style="Heading 2")
            source_heading.add_run(f"Source {source_number}: {article.title}")
            for label, value in (
                ("Date", article.published_date.isoformat()),
                ("Publisher", article.publisher),
                ("Language", article.language.upper()),
            ):
                paragraph = document.add_paragraph()
                _set_font(paragraph.add_run(f"{label}: "), size=11, bold=True)
                _set_font(paragraph.add_run(value), size=11)
            link = document.add_paragraph()
            _set_font(link.add_run("Original URL: "), size=11, bold=True)
            _add_hyperlink(link, str(article.original_url), str(article.original_url))
            if article.body is None:
                unavailable = document.add_paragraph()
                _set_font(unavailable.add_run("Full text unavailable"), size=11, bold=True, color="9B1C1C")
                continue
            summary_heading = document.add_paragraph(style="Heading 3")
            summary_heading.add_run("English Summary")
            document.add_paragraph(article.english_summary or "")
            body_heading = document.add_paragraph(style="Heading 3")
            body_heading.add_run("Article Body")
            document.add_paragraph(article.body)
    document.save(output_path)
    return output_path
