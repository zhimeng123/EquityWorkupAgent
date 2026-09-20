import json
import shutil

import fixtures
import pytest
from audit_report import run_audit
from docx import Document


def _rewrite_evidence(run, mutate):
    path = run / "evidence.json"
    payload = json.loads(path.read_text(encoding="utf-8"))
    mutate(payload)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _audit(run, template):
    return run_audit(run, template, expected_sha=fixtures.ORIGINAL_SHA256, render=False)


def test_good_run_passes(built_run, template):
    audit = _audit(built_run, template)
    assert audit.passed, audit.errors


def test_audit_detects_missing_citation(run_copy, template):
    _rewrite_evidence(run_copy, lambda p: p["fields"]["co.total_asset"].pop("source_url"))
    audit = _audit(run_copy, template)
    assert not audit.passed
    assert any("source_url" in e for e in audit.errors)


def test_audit_detects_unresolved_with_value(run_copy, template):
    _rewrite_evidence(run_copy, lambda p: p["fields"]["us.revenue"].update({"value": "USD 1"}))
    audit = _audit(run_copy, template)
    assert not audit.passed
    assert any("unresolved field must not carry a value" in e for e in audit.errors)


def test_audit_detects_invalid_status(run_copy, template):
    _rewrite_evidence(
        run_copy,
        lambda p: p["fields"]["co.soe_status"].update({"status": "guessed"}),
    )
    audit = _audit(run_copy, template)
    assert not audit.passed


def test_audit_detects_missing_gap(run_copy, template):
    (run_copy / "gaps.json").write_text(json.dumps({"gaps": []}), encoding="utf-8")
    audit = _audit(run_copy, template)
    assert not audit.passed
    assert any("gaps.json" in e for e in audit.errors)


def test_audit_detects_manual_region_change(run_copy, template):
    doc = Document(str(run_copy / "result.docx"))
    doc.tables[0].rows[0].cells[1].paragraphs[0].add_run("Hacked")
    doc.save(str(run_copy / "result.docx"))
    audit = _audit(run_copy, template)
    assert not audit.passed
    assert any("manual-only region" in e for e in audit.errors)


def test_audit_detects_third_peer_residue(run_copy, template):
    doc = Document(str(run_copy / "result.docx"))
    doc.paragraphs[0].add_run(" (peer 3)")
    doc.save(str(run_copy / "result.docx"))
    audit = _audit(run_copy, template)
    assert not audit.passed
    assert any("third competitor" in e for e in audit.errors)


def test_audit_detects_wrong_template_hash(run_copy, template):
    audit = run_audit(run_copy, template, expected_sha="0" * 64, render=False)
    assert not audit.passed
    assert any("expected SHA-256" in e for e in audit.errors)


def test_audit_detects_unwritten_writable_field(run_copy, template):
    shutil.copyfile(run_copy / "working-template.docx", run_copy / "result.docx")
    audit = _audit(run_copy, template)
    assert not audit.passed
    assert any("was not written" in e for e in audit.errors)


def test_audit_detects_duplicate_companies(run_copy, template):
    plan_path = run_copy / "run-plan.json"
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    plan["peers"][1]["ticker"] = plan["peers"][0]["ticker"]
    plan_path.write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
    audit = _audit(run_copy, template)
    assert not audit.passed
    assert any("distinct" in e for e in audit.errors)


def test_audit_detects_extra_peer_in_metrics(run_copy, template):
    metrics_path = run_copy / "metrics.json"
    metrics = json.loads(metrics_path.read_text(encoding="utf-8"))
    metrics["peers"].append({"name": "第三家竞品", "ticker": "000001"})
    metrics_path.write_text(json.dumps(metrics, ensure_ascii=False), encoding="utf-8")
    audit = _audit(run_copy, template)
    assert not audit.passed
    assert any("metrics.json peers" in e for e in audit.errors)


def test_audit_requires_peer_selection_mode(run_copy, template):
    plan_path = run_copy / "run-plan.json"
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    plan.pop("peer_selection", None)
    plan_path.write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
    audit = _audit(run_copy, template)
    assert not audit.passed
    assert any("peer_selection.mode" in e for e in audit.errors)


def test_audit_requires_authorization_for_auto_peers(run_copy, template):
    plan_path = run_copy / "run-plan.json"
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    plan["peer_selection"] = {"mode": "user_authorized_auto"}
    plan_path.write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
    audit = _audit(run_copy, template)
    assert not audit.passed
    assert any("authorization" in e for e in audit.errors)


def test_audit_accepts_authorized_auto_peers(run_copy, template):
    plan_path = run_copy / "run-plan.json"
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    plan["peer_selection"] = {
        "mode": "user_authorized_auto",
        "authorization": "另外两家竞品你自己找",
        "rationale": "memory-semiconductor comparables",
    }
    plan_path.write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
    audit = _audit(run_copy, template)
    assert audit.passed, audit.errors


def test_render_produces_pages(built_run, template):
    audit = run_audit(built_run, template, expected_sha=fixtures.ORIGINAL_SHA256, render=True)
    render = audit.checks.get("render")
    if render is None:
        pytest.skip("no host renderer available in this environment")
    assert render["pages"] > 0
    assert render["layout_issues"] == 0, audit.errors
    assert (built_run / "renders" / "page-1.png").exists()
