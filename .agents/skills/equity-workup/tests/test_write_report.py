import json

import fixtures
import pytest
from docx import Document
from docx.oxml.ns import qn
from write_report import MANUAL_FIELD_IDS, REGISTRY, bookmark_name, prepare, sha256_file, write


def _body_text(doc):
    return "".join(t.text or "" for t in doc.element.body.iter(qn("w:t")))


def _bookmarks(doc):
    return {b.get(qn("w:name")) for b in doc.element.body.iter(qn("w:bookmarkStart"))}


def _peer_rows(doc):
    table = doc.tables[7]
    nested = table.rows[1].cells[0].tables[0]
    return [[c.text.strip() for c in row.cells] for row in nested.rows]


@pytest.fixture()
def prepared(template, tmp_path):
    working = tmp_path / "working-template.docx"
    manifest = tmp_path / "manifest.json"
    result = prepare(str(template), str(working), str(manifest),
                     expected_sha256=fixtures.ORIGINAL_SHA256)
    return working, manifest, result


def test_prepare_keeps_original_untouched(template, prepared):
    _, _, result = prepared
    assert result["original_sha256"] == fixtures.ORIGINAL_SHA256
    assert sha256_file(template) == fixtures.ORIGINAL_SHA256


def test_prepare_removes_third_peer_and_wording(prepared):
    working, _, _ = prepared
    doc = Document(str(working))
    text = _body_text(doc)
    assert "(peer 3)" not in text and "peer 3" not in text
    assert "选取3家" not in text
    assert text.count("选取2家") == 2
    rows = _peer_rows(doc)
    assert [r[0] for r in rows[:4]] == ["Company Name", "Proposer", "(peer 1)", "(peer 2)"]
    assert len(rows) == 5


def test_prepare_adds_a_bookmark_for_every_registry_field(prepared):
    working, _, _ = prepared
    doc = Document(str(working))
    names = _bookmarks(doc)
    for entry in REGISTRY:
        assert bookmark_name(entry["id"]) in names
    assert len(names) == len(REGISTRY)


def test_prepare_rejects_wrong_template_hash(template, tmp_path):
    with pytest.raises(ValueError):
        prepare(str(template), str(tmp_path / "w.docx"), str(tmp_path / "m.json"),
                expected_sha256="deadbeef")


def test_write_refuses_manual_fields(prepared, tmp_path):
    working, _, _ = prepared
    evidence = tmp_path / "evidence.json"
    evidence.write_text(json.dumps({"fields": {
        "manual.underwriter": {"status": "supported", "value": "X"}
    }}), encoding="utf-8")
    with pytest.raises(ValueError, match="manual-only"):
        write(str(working), str(evidence), str(tmp_path / "out.docx"))


def test_write_refuses_unknown_fields(prepared, tmp_path):
    working, _, _ = prepared
    evidence = tmp_path / "evidence.json"
    evidence.write_text(json.dumps({"fields": {
        "co.made_up": {"status": "supported", "value": "X"}
    }}), encoding="utf-8")
    with pytest.raises(ValueError, match="unknown field id"):
        write(str(working), str(evidence), str(tmp_path / "out.docx"))


def test_write_skips_unresolved_and_writes_supported(prepared, tmp_path):
    working, _, _ = prepared
    evidence = tmp_path / "evidence.json"
    evidence.write_text(json.dumps({"fields": {
        "co.total_asset": fixtures.supported("RMB 1"),
        "us.revenue": fixtures.unresolved("not disclosed", ["path a", "path b"]),
    }}), encoding="utf-8")
    out = tmp_path / "out.docx"
    report = write(str(working), str(evidence), str(out))
    assert report["written"] == ["co.total_asset"]
    assert report["skipped"] == ["us.revenue"]

    doc = Document(str(out))
    text = _body_text(doc)
    assert "RMB 1" in text
    names = _bookmarks(doc)
    assert bookmark_name("co.total_asset") not in names
    assert bookmark_name("us.revenue") in names


def test_write_inserts_two_charts(built_run, tmp_path):
    doc = Document(str(built_run / "result.docx"))
    assert len(doc.inline_shapes) == 2
    width_emu = doc.inline_shapes[0].width
    assert width_emu == pytest.approx(int(6.3 * 914400), rel=1e-3)


def test_all_manual_ids_are_not_in_registry():
    registry_ids = {entry["id"] for entry in REGISTRY}
    assert MANUAL_FIELD_IDS.isdisjoint(registry_ids)


def test_write_never_changes_manual_regions(built_run, template):
    original = Document(str(template))
    result = Document(str(built_run / "result.docx"))
    for idx in (1, 22, 23, 24, 25, 26):
        before = "\n".join(
            " | ".join(c.text.strip() for c in row.cells) for row in original.tables[idx - 1].rows
        )
        after = "\n".join(
            " | ".join(c.text.strip() for c in row.cells) for row in result.tables[idx - 1].rows
        )
        assert before == after
