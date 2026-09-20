#!/usr/bin/env python3
"""Audit a workup run and prepare page-by-page layout review.

Checks field statuses and evidence, company identity, peer consistency, manual
regions, original-template integrity, output artifacts, the Word write result,
and — with ``--render`` — page layout via a host renderer (Microsoft Word or
LibreOffice) plus PyMuPDF. Writes ``audit.json`` and ``renders/page-N.png``.

Exit code is 0 only when every hard check passes.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import subprocess
import sys
import tempfile
from collections.abc import Sequence
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn

STATUSES = {"supported", "derived", "not_disclosed", "not_applicable", "unresolved"}
WRITABLE = {"supported", "derived", "not_disclosed", "not_applicable"}
CITATION_KEYS = (
    "source_url", "source_title", "source_type", "published_at",
    "captured_at", "page_or_section", "evidence_quote",
)
SOURCE_TYPES = {
    "annual_report", "interim_report", "quarterly_report", "exchange_filing",
    "regulator_filing", "company_official", "market_data", "media", "other_primary",
}
REQUIRED_ARTIFACTS = [
    "run-plan.json", "evidence.json", "gaps.json", "metrics.json",
    "standalone-stock-chart.png", "competitor-comparison-chart.png",
    "working-template.docx", "template-manifest.json", "result.docx",
]
MANUAL_TABLE_INDEXES = [1, 22, 23, 24, 25, 26]
MANUAL_PARA_PREFIXES = ("Sign-off:", "Date:")
MIN_CHART_PX = (1600, 750)


class Audit:
    def __init__(self) -> None:
        self.errors: list[str] = []
        self.warnings: list[str] = []
        self.checks: dict[str, Any] = {}

    def fail(self, message: str) -> None:
        self.errors.append(message)

    def warn(self, message: str) -> None:
        self.warnings.append(message)

    def ok(self, name: str, detail: Any = True) -> None:
        self.checks[name] = detail

    @property
    def passed(self) -> bool:
        return not self.errors


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _load_json(path: Path) -> Any:
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def _body_text(doc: Document) -> str:
    return "".join(t.text or "" for t in doc.element.body.iter(qn("w:t")))


def _bookmarks(doc: Document) -> set[str]:
    return {
        b.get(qn("w:name"))
        for b in doc.element.body.iter(qn("w:bookmarkStart"))
        if b.get(qn("w:name"))
    }


def _manual_signature(doc: Document) -> dict[str, str]:
    signature: dict[str, str] = {}
    for idx in MANUAL_TABLE_INDEXES:
        if idx - 1 < len(doc.tables):
            table = doc.tables[idx - 1]
            rows = []
            for row in table.rows:
                rows.append(" | ".join(c.text.strip() for c in row.cells))
            signature[f"table{idx}"] = "\n".join(rows)
    for para in doc.paragraphs:
        text = para.text.strip()
        for prefix in MANUAL_PARA_PREFIXES:
            if text.startswith(prefix):
                signature[prefix] = text
    return signature


def _peer_table(doc: Document):
    parent = doc.tables[7]
    cell = parent.rows[1].cells[0]
    return cell.tables[0] if cell.tables else None


def check_artifacts(run: Path, audit: Audit) -> None:
    missing = [name for name in REQUIRED_ARTIFACTS if not (run / name).exists()]
    if missing:
        audit.fail(f"missing artifacts: {', '.join(missing)}")
    else:
        audit.ok("artifacts", REQUIRED_ARTIFACTS)


def check_template_integrity(run: Path, template: Path, expected_sha: str | None,
                             audit: Audit) -> None:
    manifest = _load_json(run / "template-manifest.json")
    original_sha = sha256_file(template)
    if original_sha != manifest.get("original_sha256"):
        audit.fail("original template changed since prepare")
    if expected_sha and original_sha != expected_sha:
        audit.fail("original template does not match the expected SHA-256")
    working = run / "working-template.docx"
    if sha256_file(working) != manifest.get("working_sha256"):
        audit.fail("working template changed since prepare")
    audit.ok("template_integrity", {
        "original_sha256": original_sha,
        "original_unchanged": original_sha == manifest.get("original_sha256"),
    })


def check_run_plan(run: Path, audit: Audit) -> dict[str, Any]:
    plan = _load_json(run / "run-plan.json")
    target = plan.get("target") or {}
    peers = plan.get("peers") or []
    if not target.get("name") or not target.get("ticker"):
        audit.fail("run-plan target is incomplete")
    if len(peers) != 2:
        audit.fail(f"run-plan must contain exactly two competitors, found {len(peers)}")
    for peer in peers:
        if not peer.get("name") or not peer.get("ticker"):
            audit.fail("run-plan competitor is incomplete")
    tickers = [target.get("ticker")] + [p.get("ticker") for p in peers]
    if len(set(tickers)) != len(tickers):
        audit.fail("target and competitors must be distinct")
    audit.ok("run_plan", {"target": target.get("ticker"),
                          "peers": [p.get("ticker") for p in peers]})
    return plan


def check_peer_selection(plan: dict[str, Any], audit: Audit) -> None:
    selection = plan.get("peer_selection") or {}
    mode = selection.get("mode")
    if mode not in {"user_named", "user_authorized_auto"}:
        audit.fail(
            "run-plan must record peer_selection.mode "
            "(user_named or user_authorized_auto)"
        )
        return
    if mode == "user_authorized_auto":
        if not selection.get("authorization"):
            audit.fail("auto peer selection requires recorded user authorization")
        if not selection.get("rationale"):
            audit.fail("auto peer selection requires a comparability rationale")
    for peer in plan.get("peers", []):
        ticker = str(peer.get("ticker", ""))
        if not (len(ticker) == 6 and ticker.isdigit()):
            audit.fail(f"competitor ticker is not a 6-digit A-share code: {ticker!r}")
        if peer.get("exchange") not in {"SSE", "SZSE", "BSE"}:
            audit.fail(f"competitor exchange is not A-share: {peer.get('exchange')!r}")
    audit.ok("peer_selection", {"mode": mode})


def check_evidence(run: Path, audit: Audit) -> tuple[dict[str, Any], set[str]]:
    evidence = _load_json(run / "evidence.json")
    fields = evidence.get("fields", evidence)
    if not isinstance(fields, dict) or not fields:
        audit.fail("evidence.json has no fields")
        return {}, set()
    unresolved: set[str] = set()
    for field_id, record in fields.items():
        if not isinstance(record, dict):
            audit.fail(f"{field_id}: record must be an object")
            continue
        status = record.get("status")
        if status not in STATUSES:
            audit.fail(f"{field_id}: invalid status {status!r}")
            continue
        if record.get("field_id") not in (None, field_id):
            audit.fail(f"{field_id}: field_id mismatch {record.get('field_id')!r}")
        if status == "unresolved":
            unresolved.add(field_id)
            if record.get("value"):
                audit.fail(f"{field_id}: unresolved field must not carry a value")
            if not record.get("reason"):
                audit.fail(f"{field_id}: unresolved field needs a reason")
            if not record.get("attempts"):
                audit.fail(f"{field_id}: unresolved field needs an attempts list")
            continue
        if not str(record.get("value", "")).strip():
            audit.fail(f"{field_id}: written field has an empty value")
        if status == "derived":
            if not record.get("calculation"):
                audit.fail(f"{field_id}: derived field needs a calculation")
            if not record.get("derived_from"):
                audit.fail(f"{field_id}: derived field needs derived_from")
        else:
            for key in CITATION_KEYS:
                if not record.get(key):
                    audit.fail(f"{field_id}: {status} field missing {key}")
            if record.get("source_type") not in SOURCE_TYPES:
                audit.fail(f"{field_id}: invalid source_type {record.get('source_type')!r}")
    audit.ok("evidence_fields", len(fields))
    return fields, unresolved


def check_gaps(run: Path, unresolved: set[str], audit: Audit) -> None:
    gaps = _load_json(run / "gaps.json")
    entries = gaps.get("gaps", gaps)
    listed = {g.get("field_id") for g in entries} if isinstance(entries, list) else set()
    missing = unresolved - listed
    if missing:
        audit.fail(f"unresolved fields missing from gaps.json: {', '.join(sorted(missing))}")
    for gap in entries if isinstance(entries, list) else []:
        if not gap.get("reason") or not gap.get("attempts"):
            audit.fail(f"gap {gap.get('field_id')}: needs reason and attempts")
    audit.ok("gaps", len(listed))


def check_result(run: Path, template: Path, fields: dict[str, Any],
                 unresolved: set[str], plan: dict[str, Any], audit: Audit) -> None:
    result_path = run / "result.docx"
    working = Document(str(run / "working-template.docx"))
    result = Document(str(result_path))
    original = Document(str(template))

    result_text = _body_text(result)
    if "(peer 3)" in result_text or "peer 3" in result_text:
        audit.fail("third competitor placeholder remains in result.docx")
    if "选取3家" in result_text:
        audit.fail("three-peer wording remains in result.docx")
    for placeholder in ("(peer 1)", "(peer 2)"):
        if placeholder in result_text:
            audit.fail(f"unfilled competitor placeholder remains: {placeholder}")

    peer_table = _peer_table(result)
    if peer_table is None:
        audit.fail("peer comparison table missing from result.docx")
    else:
        data_rows = [r for r in peer_table.rows if any(c.text.strip() for c in r.cells)]
        if len(data_rows) != 4:
            audit.fail(
                "peer table must have a header plus three companies, "
                f"found {len(data_rows)} rows"
            )
        names = [r.cells[0].text.strip() for r in data_rows[1:]]
        target_name = (plan.get("target") or {}).get("name", "")
        peer_names = [p.get("name") for p in plan.get("peers", [])]
        if target_name and target_name not in names[0]:
            audit.fail(f"target name {target_name!r} missing from peer table")
        for name in peer_names:
            if name and name not in result_text:
                audit.fail(f"competitor {name!r} missing from result.docx")

    # Bookmark invariant: written fields lose their bookmark, untouched fields keep it.
    working_bookmarks = _bookmarks(working)
    result_bookmarks = _bookmarks(result)
    for field_id, record in fields.items():
        name = "EQW_" + field_id.replace(".", "_").replace("-", "_")
        if name not in working_bookmarks:
            continue
        status = record.get("status") if isinstance(record, dict) else None
        if status in WRITABLE:
            if name in result_bookmarks:
                audit.fail(f"{field_id}: writable field was not written")
        elif name not in result_bookmarks:
            audit.fail(f"{field_id}: non-writable field appears to have been written")

    # Manual regions must match the original template byte-for-byte in text.
    if _manual_signature(original) != _manual_signature(result):
        audit.fail("manual-only region changed in result.docx")
    audit.ok("manual_regions_unchanged", True)


def check_charts(run: Path, audit: Audit) -> None:
    for name in ("standalone-stock-chart.png", "competitor-comparison-chart.png"):
        path = run / name
        if not path.exists():
            continue
        data = path.read_bytes()
        if data[:8] != b"\x89PNG\r\n\x1a\n":
            audit.fail(f"{name}: not a PNG")
            continue
        width = int.from_bytes(data[16:20], "big")
        height = int.from_bytes(data[20:24], "big")
        if width < MIN_CHART_PX[0] or height < MIN_CHART_PX[1]:
            audit.fail(f"{name}: {width}x{height} below {MIN_CHART_PX}")
        else:
            audit.ok(f"chart_{name}", {"width": width, "height": height})


def check_identity_consistency(run: Path, plan: dict[str, Any], audit: Audit) -> None:
    """Chart/metric inputs must contain only the target and the two named peers."""
    target = plan.get("target") or {}
    peers = plan.get("peers") or []
    expected_peer_names = {p.get("name") for p in peers}
    expected_peer_tickers = {p.get("ticker") for p in peers}

    metrics = _load_json(run / "metrics.json")
    m_peers = metrics.get("peers") or []
    m_peer_names = {p.get("name") for p in m_peers}
    m_peer_tickers = {p.get("ticker") for p in m_peers}
    if (
        len(m_peers) != 2
        or m_peer_names != expected_peer_names
        or m_peer_tickers != expected_peer_tickers
    ):
        audit.fail("metrics.json peers do not match the two competitors in run-plan.json")
    m_target = metrics.get("target") or {}
    if m_target.get("ticker") != target.get("ticker"):
        audit.fail("metrics.json target does not match run-plan.json")

    evidence = _load_json(run / "evidence.json")
    if isinstance(evidence, dict) and evidence.get("peers") is not None:
        e_peers = evidence.get("peers") or []
        if len(e_peers) != 2 or {p.get("name") for p in e_peers} != expected_peer_names:
            audit.fail("evidence.json peers do not match run-plan.json")
    audit.ok("identity_consistency", {
        "target": target.get("ticker"),
        "peers": sorted(expected_peer_tickers),
    })


# --------------------------------------------------------------------------- #
# Rendering + layout
# --------------------------------------------------------------------------- #
def _soffice_path() -> str | None:
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    for candidate in (
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        "/usr/bin/soffice",
    ):
        if Path(candidate).exists():
            return candidate
    return None


def _render_with_libreoffice(docx_path: Path, out_pdf: Path) -> bool:
    soffice = _soffice_path()
    if not soffice:
        return False
    profile = Path(tempfile.mkdtemp()) / "profile"
    proc = subprocess.run(
        [soffice, "--headless", "--norestore",
         f"-env:UserInstallation=file://{profile}",
         "--convert-to", "pdf", "--outdir", str(out_pdf.parent), str(docx_path)],
        capture_output=True, text=True, check=False, timeout=300,
    )
    produced = out_pdf.parent / (docx_path.stem + ".pdf")
    if produced.exists() and produced != out_pdf:
        produced.rename(out_pdf)
    return proc.returncode == 0 and out_pdf.exists()


def _render_with_word(docx_path: Path, out_pdf: Path) -> bool:
    if not (sys.platform == "darwin" and Path("/Applications/Microsoft Word.app").exists()):
        return False
    script = f'''
with timeout of 300 seconds
set inPath to "{docx_path}"
set outPath to "{out_pdf}"
tell application "Microsoft Word"
    activate
    open inPath
    delay 4
    if (count of documents) is 0 then return "NO DOC"
    save as document 1 file name outPath file format format PDF
    delay 2
    close document 1 saving no
end tell
end timeout
return "OK"
'''
    try:
        proc = subprocess.run(["osascript", "-e", script], capture_output=True,
                              text=True, timeout=330)
    except subprocess.TimeoutExpired:
        return False
    return proc.returncode == 0 and out_pdf.exists()


def render_docx_to_pdf(docx_path: Path, out_pdf: Path) -> str | None:
    if _render_with_libreoffice(docx_path, out_pdf):
        return "libreoffice"
    if _render_with_word(docx_path, out_pdf):
        return "microsoft_word"
    return None


def _blocks_overlap(a, b, threshold: float = 0.5) -> bool:
    ax0, ay0, ax1, ay1 = a[:4]
    bx0, by0, bx1, by1 = b[:4]
    ix = max(0.0, min(ax1, bx1) - max(ax0, bx0))
    iy = max(0.0, min(ay1, by1) - max(ay0, by0))
    inter = ix * iy
    if inter < 4.0:
        return False
    smaller = min((ax1 - ax0) * (ay1 - ay0), (bx1 - bx0) * (by1 - by0))
    return smaller > 0 and inter / smaller >= threshold


def _line_boxes(page) -> list[tuple[float, float, float, float, str]]:
    """Per-line text boxes; more precise than block boxes for wrapped cells."""
    boxes: list[tuple[float, float, float, float, str]] = []
    for block in page.get_text("dict")["blocks"]:
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            text = "".join(span["text"] for span in line["spans"]).strip()
            if text:
                x0, y0, x1, y1 = line["bbox"]
                boxes.append((x0, y0, x1, y1, text))
    return boxes


def render_and_check(run: Path, audit: Audit) -> None:
    try:
        import pymupdf
    except ImportError:
        audit.warn("pymupdf not installed; page rendering skipped")
        return
    renders = run / "renders"
    renders.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        pdf = Path(tmp) / "result.pdf"
        engine = render_docx_to_pdf(run / "result.docx", pdf)
        if not engine:
            audit.warn("no Word/LibreOffice renderer available; page rendering skipped")
            return
        pdf_doc = pymupdf.open(pdf)
        page_count = pdf_doc.page_count
        issues: list[str] = []
        for index, page in enumerate(pdf_doc):
            pix = page.get_pixmap(dpi=110)
            pix.save(renders / f"page-{index + 1}.png")
            rect = page.rect
            blocks = _line_boxes(page)
            for block in blocks:
                x0, y0, x1, y1 = block[:4]
                if x0 < -1 or y0 < -1 or x1 > rect.width + 1 or y1 > rect.height + 1:
                    issues.append(f"page {index + 1}: text outside page bounds")
            for i in range(len(blocks)):
                for j in range(i + 1, len(blocks)):
                    if _blocks_overlap(blocks[i], blocks[j]):
                        issues.append(
                            f"page {index + 1}: overlapping text blocks "
                            f"{blocks[i][4].strip()[:30]!r} / {blocks[j][4].strip()[:30]!r}"
                        )
        images = sum(len(page.get_images(full=True)) for page in pdf_doc)
        pdf_doc.close()
        if issues:
            for issue in issues[:20]:
                audit.fail(issue)
        if images < 2:
            audit.warn(f"only {images} embedded image(s) found in rendered PDF")
        audit.ok("render", {"engine": engine, "pages": page_count,
                            "layout_issues": len(issues), "images": images})


# --------------------------------------------------------------------------- #
def run_audit(run: Path, template: Path, expected_sha: str | None = None,
              render: bool = False) -> Audit:
    audit = Audit()
    check_artifacts(run, audit)
    if not audit.passed:
        return audit
    check_template_integrity(run, template, expected_sha, audit)
    plan = check_run_plan(run, audit)
    check_peer_selection(plan, audit)
    fields, unresolved = check_evidence(run, audit)
    check_gaps(run, unresolved, audit)
    check_identity_consistency(run, plan, audit)
    check_result(run, template, fields, unresolved, plan, audit)
    check_charts(run, audit)
    if render:
        render_and_check(run, audit)
    return audit


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run", required=True)
    parser.add_argument("--template", required=True)
    parser.add_argument("--expected-sha256")
    parser.add_argument("--render", action="store_true")
    parser.add_argument("--json-out")
    args = parser.parse_args(argv)

    audit = run_audit(Path(args.run), Path(args.template),
                      expected_sha=args.expected_sha256, render=args.render)
    report = {
        "passed": audit.passed,
        "errors": audit.errors,
        "warnings": audit.warnings,
        "checks": audit.checks,
        "audited_at": datetime.now(UTC).isoformat(),
    }
    out = Path(args.json_out) if args.json_out else Path(args.run) / "audit.json"
    out.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"passed": audit.passed, "errors": audit.errors,
                      "warnings": audit.warnings, "audit": str(out)},
                     ensure_ascii=False, indent=2))
    return 0 if audit.passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
