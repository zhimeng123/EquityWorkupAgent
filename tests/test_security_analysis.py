from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

from docx import Document
import httpx
from pydantic import TypeAdapter

from mlc_agent.config import load_yaml
from mlc_agent.docx_writer import validate_template_mapping, write_docx_by_mapping
from mlc_agent.market_history import (
    MarketHistoryResult,
    MarketInstrument,
    MarketPoint,
    MarketSeries,
    MarketSnapshot,
    align_market_series,
    analyze_market_performance,
)
from mlc_agent.schemas import FieldMapping, FieldResult
from mlc_agent.schemas import CompanyIdentity
from mlc_agent.security_analysis import (
    IpoEvidence,
    SecuritiesOffering,
    SecuritiesOfferingReview,
    collect_security_analysis,
    select_ipo_evidence,
    select_offerings_in_window,
    security_analysis_node,
)


ROOT = Path(__file__).resolve().parents[1]
CAPTURED = datetime(2026, 7, 3, tzinfo=timezone.utc)


def _series(code, closes, *, kind):
    return MarketSeries(
        instrument=MarketInstrument(
            stock_code=code,
            display_name=f"Name {code}",
            eastmoney_secid=f"0.{code}",
            kind=kind,
        ),
        adjustment="forward_adjusted",
        points=[
            MarketPoint(trading_date=day, close=Decimal(close))
            for day, close in zip(
                [date(2024, 7, 3), date(2025, 7, 3), date(2026, 7, 3)],
                closes,
                strict=True,
            )
        ],
        source_url=f"https://market.example/{code}?fqt=1",
    )


def _market():
    aligned = align_market_series(
        _series("000938", ["100", "120", "84"], kind="target"),
        _series("399001", ["100", "110", "95"], kind="benchmark"),
        [
            _series("600001", ["100", "105", "96"], kind="peer"),
            _series("600002", ["100", "110", "97"], kind="peer"),
            _series("600003", ["100", "115", "98"], kind="peer"),
        ],
    )
    return MarketHistoryResult(
        aligned_history=aligned,
        snapshot=MarketSnapshot(
            quote_date=date(2026, 7, 3),
            current_price=Decimal("29.03"),
            week_52_high=Decimal("34.43"),
            week_52_low=Decimal("23.17"),
            source_url="https://market.example/000938?fqt=0",
        ),
        analysis=analyze_market_performance(aligned),
    )


def _offering(offering_id, announcement_date):
    return SecuritiesOffering(
        offering_id=offering_id,
        offering_type="Private placement",
        announcement_date=announcement_date,
        size_cny=Decimal("500000000"),
        price_cny=Decimal("25.50"),
        status="Completed",
        source="cninfo",
        source_url=f"https://static.cninfo.com.cn/{offering_id}.pdf",
    )


def _review(items):
    return SecuritiesOfferingReview(
        window_start=date(2025, 7, 3),
        window_end=date(2026, 7, 3),
        catalog_source="cninfo",
        catalog_source_url="https://www.cninfo.com.cn/company/000938/window",
        offerings=items,
    )


def test_offering_window_is_inclusive_and_supports_multiple_items():
    selected = select_offerings_in_window(
        [
            _offering("before", date(2025, 7, 2)),
            _offering("start", date(2025, 7, 3)),
            _offering("middle", date(2026, 1, 1)),
            _offering("end", date(2026, 7, 3)),
            _offering("after", date(2026, 7, 4)),
        ],
        as_of=date(2026, 7, 3),
    )
    assert [item.offering_id for item in selected] == ["start", "middle", "end"]


def test_security_analysis_emits_traceable_market_and_offering_fields():
    result = collect_security_analysis(
        market=_market(),
        ipo_candidates=[
            IpoEvidence(
                ipo_date=date(1999, 11, 4),
                source="exchange",
                source_url="https://www.szse.cn/000938/profile",
                evidence_period="Exchange company profile accessed 2026-07-03",
            )
        ],
        offering_review=_review(
            [_offering("offering-a", date(2025, 7, 3)), _offering("offering-b", date(2026, 7, 3))]
        ),
        as_of=date(2026, 7, 3),
        captured_at=CAPTURED,
    )
    assert result.errors == []
    by_id = {item.field_id: item for item in result.source_values}
    assert {
        "ipo_date",
        "align_to_index",
        "align_to_peers",
        "significant_stock_drop",
        "significant_stock_drop_details",
        "current_price",
        "week_52_high",
        "week_52_low",
        "securities_offerings",
    } == set(by_id)
    assert by_id["significant_stock_drop"].value == "Yes"
    assert "30.00%" in by_id["significant_stock_drop_details"].value
    assert by_id["align_to_index"].metadata["adjustment"] == "forward_adjusted"
    assert by_id["current_price"].metadata["adjustment"] == "unadjusted"
    assert len(by_id["securities_offerings"].item_evidence) == 2
    for field_id in ("align_to_index", "align_to_peers", "significant_stock_drop"):
        evidence = by_id[field_id].derived_evidence
        assert evidence and evidence.formula and evidence.inputs
        assert all(item.period and item.source_url for item in evidence.inputs)


def test_no_offering_requires_verified_complete_window_review():
    without_review = collect_security_analysis(
        market=_market(),
        ipo_candidates=[],
        offering_review=None,
        as_of=date(2026, 7, 3),
        captured_at=CAPTURED,
    )
    ids = {item.field_id for item in without_review.source_values}
    assert "securities_offerings" not in ids
    assert "ipo_date" not in ids
    assert "securities_offerings: no verified complete announcement-window review" in without_review.errors

    reviewed = collect_security_analysis(
        market=_market(),
        ipo_candidates=[],
        offering_review=_review([]),
        as_of=date(2026, 7, 3),
        captured_at=CAPTURED,
    )
    value = next(item for item in reviewed.source_values if item.field_id == "securities_offerings")
    assert value.value.startswith("No.")
    assert value.source == "cninfo"
    assert value.metadata["window_complete"] is True


def test_ipo_selection_prefers_formal_evidence_and_rejects_formal_conflict():
    eastmoney = IpoEvidence(
        ipo_date=date(1999, 11, 5),
        source="eastmoney",
        source_url="https://eastmoney.example/000938",
        evidence_period="accessed 2026-07-03",
    )
    exchange = IpoEvidence(
        ipo_date=date(1999, 11, 4),
        source="exchange",
        source_url="https://www.szse.cn/000938/profile",
        evidence_period="accessed 2026-07-03",
    )
    assert select_ipo_evidence([eastmoney, exchange]) == exchange
    conflicting_report = IpoEvidence(
        ipo_date=date(1999, 11, 3),
        source="annual_report",
        source_url="https://static.cninfo.com.cn/annual.pdf",
        evidence_period="FY2025",
    )
    try:
        select_ipo_evidence([exchange, conflicting_report])
    except ValueError as exc:
        assert "conflict" in str(exc)
    else:
        raise AssertionError("conflicting formal IPO dates were accepted")


def test_part08_mapping_matches_template_and_writes_security_fields(tmp_path):
    fragment = load_yaml(ROOT / "configs" / "fields" / "part_08.yaml")
    mappings = TypeAdapter(list[FieldMapping]).validate_python(fragment["fields"])
    normalized = [item.model_dump(mode="json") for item in mappings]
    validate_template_mapping(ROOT / "Workup_template_260617-外测版.docx", normalized)
    results = {
        item["field_id"]: FieldResult(
            field_id=item["field_id"],
            label=item["label"],
            value=(" Detail" if item["field_id"] == "significant_stock_drop_details" else f"TEST-{item['field_id']}"),
            selected_source="derived",
        )
        for item in normalized
    }
    output = tmp_path / "security.docx"
    assert write_docx_by_mapping(
        ROOT / "Workup_template_260617-外测版.docx",
        output,
        normalized,
        results,
    ) == []
    document = Document(output)
    assert "TEST-ipo_date" in document.tables[10].cell(0, 1).text
    assert "TEST-align_to_index" in document.tables[10].cell(3, 1).text
    assert "TEST-align_to_peers" in document.tables[10].cell(3, 3).text
    assert "TEST-significant_stock_drop" in document.tables[11].cell(0, 2).text
    assert "Detail" in document.tables[11].cell(1, 2).text
    assert "TEST-securities_offerings" in document.tables[11].cell(2, 2).text


def test_node_controls_market_data_failure_without_losing_existing_values():
    company = CompanyIdentity(
        company_name="Target",
        company_short_name="Target",
        stock_code="000938",
        exchange="深圳证券交易所",
        eastmoney_secid="0.000938",
        eastmoney_secu_code="000938.SZ",
        xueqiu_symbol="SZ000938",
    )
    peers = [
        CompanyIdentity(
            company_name=f"Peer {code}",
            company_short_name=f"Peer {code}",
            stock_code=code,
            exchange="上海证券交易所",
            eastmoney_secid=f"1.{code}",
            eastmoney_secu_code=f"{code}.SH",
            xueqiu_symbol=f"SH{code}",
        )
        for code in ("600001", "600002", "600003")
    ]
    existing = {"field_id": "existing", "value": "keep"}
    state = {
        "company": company.model_dump(mode="json"),
        "created_at": "2026-07-03T00:00:00+00:00",
        "source_values": [existing],
        "node_errors": [{"node": "earlier", "message": "keep"}],
        "part_results": {
            "part_05": {
                "selection": {
                    "selected": [{"stock_code": item.stock_code} for item in peers]
                }
            },
            "part_08_input": {
                "as_of": "2026-07-03",
                "peer_identities": [item.model_dump(mode="json") for item in peers],
            },
        },
    }

    def handler(request):
        return httpx.Response(200, json={"data": {"klines": []}}, request=request)

    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        result = security_analysis_node(state, client=client)

    assert result["source_values"] == [existing]
    assert result["part_results"]["part_08"] == {
        "status": "failed",
        "reason": "market_history: Tencent market history returned an invalid result for 000938",
    }
    assert result["node_errors"][-1] == {
        "node": "part_08",
        "message": "market_history: Tencent market history returned an invalid result for 000938",
    }
