from __future__ import annotations

from datetime import date, datetime, timezone
import json
from pathlib import Path

import pytest
from docx import Document

from mlc_agent.company_supplement import (
    EvidenceDocument,
    collect_company_supplement,
    company_supplement_node,
    make_openai_extractor,
)
from mlc_agent.config import load_template_mapping
from mlc_agent.docx_writer import validate_template_mapping
from mlc_agent.docx_writer import write_docx_by_mapping


AS_OF = date(2026, 7, 3)
CAPTURED_AT = datetime(2026, 7, 3, 9, 0, tzinfo=timezone.utc)
ANNUAL_URL = "https://static.cninfo.com.cn/annual.pdf"
NOTICE_URL = "https://static.cninfo.com.cn/notice.pdf"
ANNUAL_TEXT = (
    "最终实际控制人为某市国资委。"
    "公司控股的上市子公司为紫光股份，股票代码000938，深圳证券交易所上市。"
    "第二上市子公司股票代码600001，在上海证券交易所上市。"
    "现任董事张三兼任甲上市公司董事。"
    "现任董事李四兼任乙上市公司独立董事。"
    "公司明确说明除上述公司外无其他上市子公司。"
    "公司未来十二个月无重大并购计划。"
)
NOTICE_TEXT = (
    "2025年7月3日，公司披露重大资产重组，交易标的为目标甲公司，位于中国，"
    "未上市，交易对价人民币10亿元，状态为进行中。"
    "2026年1月2日，公司披露重大收购目标乙公司，位于德国，已上市，"
    "交易对价欧元2亿元，状态为已完成。"
)


@pytest.fixture
def documents() -> list[EvidenceDocument]:
    return [
        EvidenceDocument(
            source="annual_report",
            source_url=ANNUAL_URL,
            disclosure_date=date(2026, 4, 20),
            text=ANNUAL_TEXT,
        ),
        EvidenceDocument(
            source="cninfo",
            source_url=NOTICE_URL,
            disclosure_date=date(2025, 7, 3),
            text=NOTICE_TEXT,
        ),
    ]


def evidence(url: str, disclosure_date: str, text: str) -> dict[str, object]:
    return {
        "source_url": url,
        "disclosure_date": disclosure_date,
        "page_number": 1,
        "evidence_text": text,
    }


def extraction(*, soe: str = "state_owned") -> dict[str, object]:
    return {
        "soe": {
            "classification": soe,
            "ultimate_controller": "某市国资委",
            "evidence": evidence(ANNUAL_URL, "2026-04-20", "最终实际控制人为某市国资委。"),
        },
        "listed_subsidiaries": {
            "status": "yes",
            "items": [
                {
                    "name": "紫光股份",
                    "stock_code": "000938",
                    "exchange": "深圳证券交易所",
                    "listing_status": "Listed",
                    "evidence": evidence(
                        ANNUAL_URL,
                        "2026-04-20",
                        "公司控股的上市子公司为紫光股份，股票代码000938，深圳证券交易所上市。",
                    ),
                },
                {
                    "name": "第二上市子公司",
                    "stock_code": "600001",
                    "exchange": "上海证券交易所",
                    "listing_status": "Listed",
                    "evidence": evidence(
                        ANNUAL_URL,
                        "2026-04-20",
                        "第二上市子公司股票代码600001，在上海证券交易所上市。",
                    ),
                },
            ],
        },
        "outside_directorships": {
            "status": "yes",
            "items": [
                {
                    "director_name": "张三",
                    "is_current_director": True,
                    "listed_company_name": "甲上市公司",
                    "position": "Director",
                    "stock_code": "600002",
                    "exchange": "上海证券交易所",
                    "evidence": evidence(
                        ANNUAL_URL, "2026-04-20", "现任董事张三兼任甲上市公司董事。"
                    ),
                },
                {
                    "director_name": "李四",
                    "is_current_director": True,
                    "listed_company_name": "乙上市公司",
                    "position": "Independent Director",
                    "evidence": evidence(
                        ANNUAL_URL, "2026-04-20", "现任董事李四兼任乙上市公司独立董事。"
                    ),
                },
            ],
        },
        "major_ma_past_12m": {
            "status": "yes",
            "items": [
                {
                    "announcement_date": "2025-07-03",
                    "transaction_target": "目标甲公司",
                    "country": "中国",
                    "target_listed": "No",
                    "consideration": "人民币10亿元",
                    "status": "进行中",
                    "evidence": evidence(
                        NOTICE_URL,
                        "2025-07-03",
                        "2025年7月3日，公司披露重大资产重组，交易标的为目标甲公司，位于中国，未上市，交易对价人民币10亿元，状态为进行中。",
                    ),
                },
                {
                    "announcement_date": "2026-01-02",
                    "transaction_target": "目标乙公司",
                    "country": "德国",
                    "target_listed": "Yes",
                    "consideration": "欧元2亿元",
                    "status": "已完成",
                    "evidence": evidence(
                        NOTICE_URL,
                        "2025-07-03",
                        "2026年1月2日，公司披露重大收购目标乙公司，位于德国，已上市，交易对价欧元2亿元，状态为已完成。",
                    ),
                },
            ],
        },
        "ma_plan_next_12m": {
            "status": "no",
            "items": [],
            "negative_evidence": evidence(
                ANNUAL_URL, "2026-04-20", "公司未来十二个月无重大并购计划。"
            ),
        },
    }


def test_deepseek_extractor_uses_plain_json_instruction_and_local_validation(documents):
    class FakeCompletions:
        def create(self, **kwargs):
            assert "response_format" not in kwargs
            assert "JSON Schema" in kwargs["messages"][0]["content"]
            request = json.loads(kwargs["messages"][1]["content"])
            message = type("Message", (), {"content": json.dumps(extraction()[request["field"]])})()
            return type("Response", (), {"choices": [type("Choice", (), {"message": message})()]})()

    client = type("Client", (), {"chat": type("Chat", (), {"completions": FakeCompletions()})()})()

    result = make_openai_extractor(client, model="deepseek-chat")(documents, as_of=AS_OF)

    assert result["soe"]["classification"] == "state_owned"


def test_deepseek_extractor_isolates_malformed_field_json(documents):
    class FakeCompletions:
        def create(self, **kwargs):
            field = json.loads(kwargs["messages"][1]["content"])["field"]
            content = "{invalid" if field == "outside_directorships" else json.dumps(extraction()[field])
            message = type("Message", (), {"content": content})()
            return type("Response", (), {"choices": [type("Choice", (), {"message": message})()]})()

    client = type("Client", (), {"chat": type("Chat", (), {"completions": FakeCompletions()})()})()
    extracted = make_openai_extractor(client, model="deepseek-chat")(documents, as_of=AS_OF)
    result = collect_company_supplement(
        documents, as_of=AS_OF, extractor=lambda _documents, as_of: extracted,
        captured_at=CAPTURED_AT,
    )

    assert {item.field_id for item in result.source_values} == {
        "soe_classification", "listed_subsidiaries", "major_ma_past_12m", "ma_plan_next_12m",
    }
    assert [item.field_id for item in result.errors] == ["listed_outside_directorships"]


@pytest.mark.parametrize(
    ("classification", "expected"),
    [("state_owned", "SOE"), ("non_state_owned", "Non-SOE")],
)
def test_soe_uses_ultimate_controller_classification(documents, classification, expected):
    result = collect_company_supplement(
        documents,
        as_of=AS_OF,
        extractor=lambda _documents, as_of: extraction(soe=classification),
        captured_at=CAPTURED_AT,
    )
    assert not result.errors
    value = next(item for item in result.source_values if item.field_id == "soe_classification")
    assert value.value.startswith(expected + " |")
    assert "某市国资委" in value.value


def test_multiple_items_keep_per_item_evidence_and_formatting(documents):
    result = collect_company_supplement(
        documents,
        as_of=AS_OF,
        extractor=lambda _documents, as_of: extraction(),
        captured_at=CAPTURED_AT,
    )
    subsidiaries = next(item for item in result.source_values if item.field_id == "listed_subsidiaries")
    directorships = next(
        item for item in result.source_values if item.field_id == "listed_outside_directorships"
    )
    transaction = next(item for item in result.source_values if item.field_id == "major_ma_past_12m")
    assert subsidiaries.value.count("\n") == 1
    assert subsidiaries.value.splitlines()[0].startswith("Listed | 紫光股份 | 000938")
    assert subsidiaries.raw_value["items"][0]["listing_status"] == "Listed"
    assert subsidiaries.item_evidence[0].raw_value["listing_status"] == "Listed"
    assert len(subsidiaries.item_evidence) == 2
    assert directorships.value.count("\n") == 1
    assert len(directorships.item_evidence) == 2
    assert "Consideration: 人民币10亿元" in transaction.value
    assert len(transaction.item_evidence) == 2
    assert all(item.source_url == NOTICE_URL for item in transaction.item_evidence)


def test_not_disclosed_is_verified_value_while_explicit_negative_is_no(documents):
    payload = extraction()
    payload["listed_subsidiaries"] = {
        "status": "not_disclosed",
        "items": [],
        "negative_evidence": evidence(
            ANNUAL_URL,
            "2026-04-20",
            "公司明确说明除上述公司外无其他上市子公司。",
        ),
    }
    result = collect_company_supplement(
        documents,
        as_of=AS_OF,
        extractor=lambda _documents, as_of: payload,
        captured_at=CAPTURED_AT,
    )
    subsidiaries = next(
        item for item in result.source_values if item.field_id == "listed_subsidiaries"
    )
    assert subsidiaries.value == "Not disclosed"
    assert not any(item.field_id == "listed_subsidiaries" for item in result.errors)
    plan = next(item for item in result.source_values if item.field_id == "ma_plan_next_12m")
    assert plan.value == "No"


@pytest.mark.parametrize("announcement_date", ["2025-07-03", "2026-07-03"])
def test_past_12_month_window_is_inclusive_at_both_boundaries(documents, announcement_date):
    payload = extraction()
    payload["major_ma_past_12m"]["items"][0]["announcement_date"] = announcement_date
    result = collect_company_supplement(
        documents,
        as_of=AS_OF,
        extractor=lambda _documents, as_of: payload,
        captured_at=CAPTURED_AT,
    )
    assert not result.errors


def test_field_validation_failure_does_not_discard_other_fields(documents):
    payload = extraction()
    payload["major_ma_past_12m"]["items"][0]["announcement_date"] = "2025-07-02"
    result = collect_company_supplement(
        documents,
        as_of=AS_OF,
        extractor=lambda _documents, as_of: payload,
        captured_at=CAPTURED_AT,
    )
    assert len(result.source_values) == 4
    assert len(result.errors) == 1
    assert result.errors[0].field_id == "major_ma_past_12m"
    assert "outside" in result.errors[0].reason

    payload = extraction()
    payload["soe"]["evidence"]["evidence_text"] = "不存在的证据"
    result = collect_company_supplement(
        documents,
        as_of=AS_OF,
        extractor=lambda _documents, as_of: payload,
        captured_at=CAPTURED_AT,
    )
    assert len(result.errors) == 1
    assert result.errors[0].field_id == "soe_classification"
    assert "could not be anchored" in result.errors[0].reason
    assert len(result.source_values) == 4


def test_missing_listed_subsidiary_negative_evidence_fails_only_that_field(documents):
    payload = extraction()
    payload["listed_subsidiaries"] = {
        "status": "no",
        "items": [],
        "negative_evidence": None,
    }

    result = collect_company_supplement(
        documents,
        as_of=AS_OF,
        extractor=lambda _documents, as_of: payload,
        captured_at=CAPTURED_AT,
    )

    assert len(result.source_values) == 4
    assert [item.field_id for item in result.errors] == ["listed_subsidiaries"]
    assert "explicit negative evidence" in result.errors[0].reason


def test_major_subsidiary_table_does_not_prove_separate_listing(documents):
    payload = extraction()
    payload["listed_subsidiaries"] = {
        "status": "yes",
        "items": [{
            "name": "紫光股份",
            "stock_code": "000938",
            "exchange": "深圳证券交易所",
            "listing_status": "Listed",
            "evidence": evidence(
                ANNUAL_URL,
                "2026-04-20",
                "公司明确说明除上述公司外无其他上市子公司。",
            ),
        }],
    }

    result = collect_company_supplement(
        documents,
        as_of=AS_OF,
        extractor=lambda _documents, as_of: payload,
        captured_at=CAPTURED_AT,
    )

    subsidiaries = next(
        item for item in result.source_values if item.field_id == "listed_subsidiaries"
    )
    assert subsidiaries.value == "Not disclosed"
    assert subsidiaries.raw_value["status"] == "not_disclosed"
    assert not any(item.field_id == "listed_subsidiaries" for item in result.errors)


def test_node_returns_part_result_without_writing_word(documents):
    state = {"part_01_documents": [item.model_dump(mode="json") for item in documents]}
    update = company_supplement_node(
        state, extractor=lambda _documents, as_of: extraction(), as_of=AS_OF
    )
    assert set(update) == {"part_results", "source_values", "node_errors"}
    assert update["part_results"]["part_01"]["raw_result"] is not None


def test_non_current_director_is_rejected_by_schema(documents):
    payload = extraction()
    payload["outside_directorships"]["items"][0]["is_current_director"] = False
    result = collect_company_supplement(
        documents,
        as_of=AS_OF,
        extractor=lambda _documents, as_of: payload,
        captured_at=CAPTURED_AT,
    )
    assert len(result.source_values) == 4
    assert len(result.errors) == 1
    assert result.errors[0].field_id == "listed_outside_directorships"
    assert "is_current_director" in result.errors[0].reason


def test_invalid_soe_does_not_discard_other_company_supplement_fields(documents):
    data = extraction()
    data["soe"] = {"classification": "non_state_owned"}

    result = collect_company_supplement(
        documents,
        as_of=AS_OF,
        extractor=lambda _documents, as_of: data,
        captured_at=CAPTURED_AT,
    )

    assert len(result.source_values) == 4
    assert [item.field_id for item in result.errors] == ["soe_classification"]
    assert "controller and evidence" in result.errors[0].reason


def test_part_01_mapping_matches_real_template():
    root = Path(__file__).resolve().parents[1]
    mapping = load_template_mapping(root / "configs")
    part_fields = [item for item in mapping["fields"] if item["field_id"] in {
        "soe_classification",
        "listed_subsidiaries",
        "listed_outside_directorships",
        "major_ma_past_12m",
        "ma_plan_next_12m",
    }]
    assert len(part_fields) == 5
    validate_template_mapping(root / "Workup_template_260617-外测版.docx", part_fields)


def test_missing_part_results_leave_original_template_targets_unchanged(tmp_path):
    root = Path(__file__).resolve().parents[1]
    mapping = load_template_mapping(root / "configs")
    part_fields = [item for item in mapping["fields"] if item["field_id"] in {
        "soe_classification",
        "listed_subsidiaries",
        "listed_outside_directorships",
        "major_ma_past_12m",
        "ma_plan_next_12m",
    }]
    output = tmp_path / "unchanged.docx"
    assert write_docx_by_mapping(
        root / "Workup_template_260617-外测版.docx", output, part_fields, {}
    ) == []
    document = Document(output)
    assert document.tables[1].cell(0, 3).text == ""
    assert "可在最新年报和半年报上抓取" in document.tables[1].cell(5, 3).text
    assert "可在最新年报和半年报上抓取" in document.paragraphs[5].text
