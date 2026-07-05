from __future__ import annotations

from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

from docx import Document
from pydantic import TypeAdapter

from mlc_agent.config import load_yaml
from mlc_agent.docx_writer import validate_template_mapping, write_docx_by_mapping
from mlc_agent.operating_performance import FinancialPeriodRecord
from mlc_agent.peer_analysis import (
    DisclosedMetric,
    PeerCompany,
    collect_peer_analysis,
    select_peers,
)
from mlc_agent.related_parties import (
    RelatedPartyTransaction,
    collect_related_party_transactions,
)
from mlc_agent.schemas import FieldMapping, FieldResult


ROOT = Path(__file__).resolve().parents[1]
CAPTURED_AT = datetime(2026, 7, 3, tzinfo=timezone.utc)


def _record(year: int, revenue: str, profit: str, gross_profit: str) -> FinancialPeriodRecord:
    return FinancialPeriodRecord(
        report_date=date(year, 12, 31),
        period_type="annual",
        fiscal_year=year,
        revenue=Decimal(revenue),
        parent_net_profit=Decimal(profit),
        gross_profit=Decimal(gross_profit),
        source_url=f"https://example.com/{year}/annual",
    )


def _company(
    code: str,
    latest_revenue: str,
    *,
    industry: str = "C39",
    a_share: bool = True,
    st: bool = False,
    financial: bool = False,
    include_previous: bool = True,
    inventory_period: str = "FY2025",
) -> PeerCompany:
    records = [
        _record(2025, latest_revenue, str(Decimal(latest_revenue) / 10), str(Decimal(latest_revenue) / 5))
    ]
    if include_previous:
        records.append(
            _record(2024, str(Decimal(latest_revenue) * Decimal("0.8")), str(Decimal(latest_revenue) / 12), str(Decimal(latest_revenue) / 6))
        )
    return PeerCompany(
        stock_code=code,
        company_name=f"Company {code}",
        csrc_industry=industry,
        is_a_share=a_share,
        is_st=st,
        is_financial=financial,
        annual_records=records,
        inventory_turnover=DisclosedMetric(
            value=Decimal("3.25"),
            period=inventory_period,
            source_url=f"https://eastmoney.example/{code}/inventory",
        ),
    )


def test_peer_selector_applies_all_exclusions_and_stable_revenue_sort():
    target = _company("000938", "100000000")
    candidates = [
        _company("000938", "100000000"),
        _company("600001", "101000000", industry="C40"),
        _company("600002", "101000000", st=True),
        _company("600003", "101000000", financial=True),
        _company("600004", "101000000", a_share=False),
        _company("600005", "101000000", include_previous=False),
        _company("600006", "101000000", inventory_period="FY2024"),
        _company("600020", "110000000"),
        _company("600010", "110000000"),
        _company("600030", "105000000"),
        _company("600040", "130000000"),
    ]
    result = select_peers(target, candidates)
    assert result.success is True
    assert [item.stock_code for item in result.selected] == ["600030", "600010", "600020"]
    exclusions = {item.stock_code: item.reason for item in result.exclusions}
    assert exclusions["000938"] == "target company"
    assert exclusions["600001"] == "different CSRC industry"
    assert exclusions["600002"] == "ST or *ST company"
    assert exclusions["600003"] == "financial company"
    assert exclusions["600004"] == "not an A-share company"
    assert "two common" in exclusions["600005"]
    assert exclusions["600006"] == "inventory turnover period mismatch"


def test_fewer_than_three_peers_fails_without_cross_industry_backfill():
    result = select_peers(
        _company("000938", "100000000"),
        [
            _company("600001", "101000000"),
            _company("600002", "102000000"),
            _company("600003", "100500000", industry="C40"),
        ],
    )
    assert result.success is False
    assert result.selected == []
    assert result.failure_reason == "Fewer than three eligible same-industry A-share peers"


def test_non_cny_candidate_is_excluded_without_conversion_fallback():
    candidate = _company("600001", "101000000")
    candidate.annual_records[0].currency = "USD"
    result = select_peers(
        _company("000938", "100000000"),
        [
            candidate,
            _company("600002", "102000000"),
            _company("600003", "103000000"),
            _company("600004", "104000000"),
        ],
    )
    assert result.success is True
    assert [item.stock_code for item in result.selected] == ["600002", "600003", "600004"]
    assert next(item.reason for item in result.exclusions if item.stock_code == "600001") == (
        "annual financial periods are not in CNY"
    )


def test_candidate_missing_inventory_turnover_is_excluded_and_next_peer_is_selected():
    missing = _company("600001", "101000000")
    missing.inventory_turnover = None
    result = select_peers(
        _company("000938", "100000000"),
        [
            missing,
            _company("600002", "102000000"),
            _company("600003", "103000000"),
            _company("600004", "104000000"),
        ],
    )
    assert result.success is True
    assert [item.stock_code for item in result.selected] == ["600002", "600003", "600004"]
    assert next(item.reason for item in result.exclusions if item.stock_code == "600001") == (
        "missing disclosed inventory turnover"
    )


def test_peer_metrics_use_fixed_net_margin_and_trace_every_metric():
    target = _company("000938", "100000000")
    result = collect_peer_analysis(
        target=target,
        candidates=[
            _company("600001", "101000000"),
            _company("600002", "102000000"),
            _company("600003", "103000000"),
        ],
        captured_at=CAPTURED_AT,
    )
    assert result.failures == []
    assert len(result.metrics) == 4
    proposer = result.metrics[0]
    assert proposer.net_margin == Decimal("10.00")
    assert proposer.gross_margin == Decimal("20.00")
    assert proposer.revenue_growth == Decimal("25.00")
    table = next(item for item in result.source_values if item.field_id == "peer_comparison_table")
    assert len(table.metadata["structured_value"]) == 4
    assert len(table.item_evidence) == 4
    assert all(item.period == "FY2025" for item in table.item_evidence)
    input_ids = {item.field_id for item in table.derived_evidence.inputs}
    for code in ("000938", "600001", "600002", "600003"):
        assert f"{code}.revenue" in input_ids
        assert f"{code}.previous_revenue" in input_ids
        assert f"{code}.gross_profit" in input_ids
        assert f"{code}.parent_net_profit" in input_ids
        assert f"{code}.inventory_turnover" in input_ids


def test_missing_peer_metric_fails_whole_peer_fields():
    target = _company("000938", "100000000")
    target.annual_records[0].gross_profit = None
    result = collect_peer_analysis(
        target=target,
        candidates=[_company("600001", "101000000"), _company("600002", "102000000"), _company("600003", "103000000")],
        captured_at=CAPTURED_AT,
    )
    assert result.source_values == []
    assert result.metrics == []
    assert result.failures == ["000938: missing gross profit"]


def test_target_missing_inventory_turnover_returns_explicit_peer_failure():
    target = _company("000938", "100000000")
    target.inventory_turnover = None
    result = collect_peer_analysis(
        target=target,
        candidates=[
            _company("600001", "101000000"),
            _company("600002", "102000000"),
            _company("600003", "103000000"),
        ],
        captured_at=CAPTURED_AT,
    )
    assert result.source_values == []
    assert result.metrics == []
    assert result.failures == ["000938: missing disclosed inventory turnover"]


def test_related_transactions_preserve_each_source_and_do_not_infer_approval():
    value = collect_related_party_transactions(
        [
            RelatedPartyTransaction(
                related_party="关联方甲",
                relationship="Controlling shareholder affiliate",
                transaction_type="Sale of goods",
                amount_cny=Decimal("12000000"),
                commercial_terms="Market pricing",
                board_approved=True,
                period="FY2025",
                source="annual_report",
                source_url="https://example.com/annual.pdf",
                source_page=120,
            ),
            RelatedPartyTransaction(
                related_party="关联方乙",
                relationship="Associate",
                transaction_type="Service purchase",
                period="2026H1",
                source="interim_report",
                source_url="https://example.com/interim.pdf",
                source_page=66,
            ),
        ],
        captured_at=CAPTURED_AT,
    )
    assert value is not None
    assert "Board approval: Not disclosed" in value.value
    assert "Terms: Not disclosed" in value.value
    assert [item.source_url for item in value.item_evidence] == [
        "https://example.com/annual.pdf",
        "https://example.com/interim.pdf",
    ]


def test_peer_nested_table_mapping_writes_proposer_and_three_peers(tmp_path):
    fragment = load_yaml(ROOT / "configs" / "fields" / "part_05.yaml")
    mappings = TypeAdapter(list[FieldMapping]).validate_python(fragment["fields"])
    normalized = [item.model_dump(mode="json") for item in mappings]
    template = ROOT / "Workup_template_260617-外测版.docx"
    validate_template_mapping(template, normalized)
    table_mapping = next(item for item in normalized if item["field_id"] == "peer_comparison_table")
    rows = [
        ["Proposer", "CNY 100.00 million", "3.00", "20.00%", "10.00%"],
        ["Peer 1", "CNY 101.00 million", "3.10", "19.00%", "9.00%"],
        ["Peer 2", "CNY 102.00 million", "3.20", "18.00%", "8.00%"],
        ["Peer 3", "CNY 103.00 million", "3.30", "17.00%", "7.00%"],
    ]
    output = tmp_path / "peer.docx"
    failures = write_docx_by_mapping(
        template,
        output,
        [table_mapping],
        {
            "peer_comparison_table": FieldResult(
                field_id="peer_comparison_table",
                label="Peer Comparison Table",
                value="table",
                selected_source="derived",
                structured_value=rows,
            )
        },
    )
    assert failures == []
    nested = Document(output).tables[7].cell(1, 0).tables[0]
    assert [nested.cell(row, 0).text for row in range(1, 5)] == [
        "Proposer",
        "Peer 1",
        "Peer 2",
        "Peer 3",
    ]
    assert nested.cell(0, 0).text == "Company Name"
