from __future__ import annotations

import base64
import builtins
from importlib.metadata import version
from datetime import date, datetime
from pathlib import Path

import pytest
import httpx
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import RGBColor

from mlc_agent.cninfo import (
    AnnouncementDocument,
    build_report_catalog,
    fetch_company_announcements,
    select_reports,
)
from mlc_agent.config import load_template_mapping
from mlc_agent.docx_writer import validate_template_mapping, write_docx_by_mapping
from mlc_agent.exceptions import DataSourceError, TemplateMappingError
from mlc_agent.report_parser import parse_machine_generated_pdf
from mlc_agent.nodes import self_check_node
from mlc_agent.schemas import FieldResult




def _machine_generated_pdf() -> bytes:
    def content(title: str, headers: tuple[str, str], values: tuple[str, str]) -> bytes:
        commands = [
            f"BT /F1 12 Tf 72 720 Td ({title}) Tj ET",
            "72 650 m 432 650 l S",
            "72 626 m 432 626 l S",
            "72 602 m 432 602 l S",
            "72 650 m 72 602 l S",
            "252 650 m 252 602 l S",
            "432 650 m 432 602 l S",
            f"BT /F1 12 Tf 78 633 Td ({headers[0]}) Tj ET",
            f"BT /F1 12 Tf 258 633 Td ({headers[1]}) Tj ET",
            f"BT /F1 12 Tf 78 609 Td ({values[0]}) Tj ET",
            f"BT /F1 12 Tf 258 609 Td ({values[1]}) Tj ET",
        ]
        return ("\n".join(commands) + "\n").encode("ascii")

    page_one = content("Annual Report Page One", ("Metric", "Value"), ("Revenue", "100"))
    page_two = content("Annual Report Page Two", ("Item", "Amount"), ("Cash", "80"))
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [4 0 R 6 0 R] /Count 2 >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents 5 0 R >>",
        b"<< /Length " + str(len(page_one)).encode("ascii") + b" >>\nstream\n" + page_one + b"endstream",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents 7 0 R >>",
        b"<< /Length " + str(len(page_two)).encode("ascii") + b" >>\nstream\n" + page_two + b"endstream",
    ]
    pdf = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{number} 0 obj\n".encode("ascii") + obj + b"\nendobj\n")
    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    return bytes(pdf)


def _result(field_id: str, value: str, **kwargs) -> FieldResult:
    return FieldResult(
        field_id=field_id,
        label=field_id,
        value=value,
        selected_source="fixed",
        **kwargs,
    )


def test_mapping_includes_fragments_and_rejects_conflicting_targets(tmp_path):
    fields_dir = tmp_path / "fields"
    fields_dir.mkdir()
    (tmp_path / "fixed_template_mapping.yaml").write_text(
        "template_version: v1\ntemplate_filename: t.docx\nincludes:\n  - fields/a.yaml\nfields: []\n",
        encoding="utf-8",
    )
    duplicate_target = """fields:
  - field_id: a
    label: A
    source_type: fixed
    locators:
      - kind: paragraph
        paragraph_index: 0
        expected_text: X
  - field_id: b
    label: B
    source_type: fixed
    locators:
      - kind: paragraph
        paragraph_index: 0
        expected_text: X
"""
    (fields_dir / "a.yaml").write_text(duplicate_target, encoding="utf-8")
    with pytest.raises(TemplateMappingError, match="conflicting locator target"):
        load_template_mapping(tmp_path)


def test_writer_supports_nested_multi_target_multiline_table_hyperlink_and_image(tmp_path):
    template = tmp_path / "template.docx"
    output = tmp_path / "result.docx"
    image = tmp_path / "pixel.png"
    image.write_bytes(
        base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
        )
    )
    document = Document()
    document.add_paragraph("PRIMARY")
    document.add_paragraph("SECONDARY")
    document.add_paragraph("LINES")
    document.add_paragraph("Link: URL")
    document.add_paragraph("Label: SUFFIX")
    top = document.add_table(rows=2, cols=2)
    top.cell(0, 0).text = "TABLE"
    nested = top.cell(1, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "NESTED"
    top.cell(1, 1).text = ""
    document.save(template)

    mappings = [
        {
            "field_id": "multi",
            "label": "Multi",
            "source_type": "fixed",
            "locators": [
                {"kind": "paragraph", "paragraph_index": 0, "expected_text": "PRIMARY"},
                {"kind": "paragraph", "paragraph_index": 1, "expected_text": "SECONDARY"},
            ],
            "write_strategy": "replace_text",
            "output_format": "text",
        },
        {
            "field_id": "lines",
            "label": "Lines",
            "source_type": "fixed",
            "locators": [{"kind": "paragraph", "paragraph_index": 2, "expected_text": "LINES"}],
            "write_strategy": "replace_multiline_text",
            "output_format": "multiline",
        },
        {
            "field_id": "link",
            "label": "Link",
            "source_type": "fixed",
            "locators": [{"kind": "paragraph", "paragraph_index": 3, "expected_text": "URL"}],
            "write_strategy": "replace_with_hyperlink",
            "output_format": "url",
        },
        {
            "field_id": "append",
            "label": "Append",
            "source_type": "fixed",
            "locators": [{"kind": "paragraph", "paragraph_index": 4, "expected_text": "Label: "}],
            "write_strategy": "append_after_label",
            "output_format": "text",
        },
        {
            "field_id": "table",
            "label": "Table",
            "source_type": "fixed",
            "locators": [{
                "kind": "table_cell",
                "table_path": [{"table_index": 0, "row_index": 0, "column_index": 0}],
                "expected_text": "TABLE",
            }],
            "write_strategy": "fill_fixed_table",
            "output_format": "table",
        },
        {
            "field_id": "nested",
            "label": "Nested",
            "source_type": "fixed",
            "locators": [{
                "kind": "table_cell",
                "table_path": [
                    {"table_index": 0, "row_index": 1, "column_index": 0},
                    {"table_index": 0, "row_index": 0, "column_index": 0},
                ],
                "expected_text": "NESTED",
            }],
            "write_strategy": "replace_text",
            "output_format": "text",
        },
        {
            "field_id": "image",
            "label": "Image",
            "source_type": "fixed",
            "locators": [{
                "kind": "image_anchor",
                "table_path": [{"table_index": 0, "row_index": 1, "column_index": 1}],
                "expected_empty": True,
            }],
            "write_strategy": "insert_image",
            "output_format": "image",
        },
    ]
    validate_template_mapping(template, mappings)
    results = {
        "multi": _result("multi", "SAME"),
        "lines": _result("lines", "one\ntwo"),
        "link": _result("link", "https://example.com"),
        "append": _result("append", "VALUE"),
        "table": _result("table", "", structured_value=[["A", "B"]]),
        "nested": _result("nested", "DONE"),
        "image": _result("image", "chart", artifact_path=str(image)),
    }
    assert write_docx_by_mapping(template, output, mappings, results) == []
    written = Document(output)
    assert written.paragraphs[0].text == written.paragraphs[1].text == "SAME"
    assert written.paragraphs[2].text == "one\ntwo"
    assert written.paragraphs[4].text == "Label: VALUESUFFIX"
    assert written.tables[0].cell(0, 0).text == "A"
    assert written.tables[0].cell(0, 1).text == "B"
    assert "DONE" in written.tables[0].cell(1, 0).tables[0].cell(0, 0).text
    assert any(rel.reltype == RT.IMAGE for rel in written.part.rels.values())
    assert any(rel.reltype == RT.HYPERLINK for rel in written.part.rels.values())


def test_failed_multi_target_rolls_back_all_targets(tmp_path):
    template = tmp_path / "template.docx"
    output = tmp_path / "result.docx"
    document = Document()
    document.add_paragraph("FIRST")
    document.add_paragraph("SECOND")
    document.save(template)
    mapping = [{
        "field_id": "field",
        "label": "Field",
        "source_type": "fixed",
        "locators": [
            {"kind": "paragraph", "paragraph_index": 0, "expected_text": "FIRST"},
            {"kind": "paragraph", "paragraph_index": 1, "expected_text": "MISSING"},
        ],
        "write_strategy": "replace_text",
        "output_format": "text",
    }]
    failures = write_docx_by_mapping(template, output, mapping, {"field": _result("field", "VALUE")})
    assert failures and failures[0][0] == "field"
    written = Document(output)
    assert written.paragraphs[0].text == "FIRST"
    assert written.paragraphs[1].text == "SECOND"


def test_mapped_failure_and_out_of_scope_blue_note_are_separated(tmp_path):
    template = tmp_path / "template.docx"
    output = tmp_path / "result.docx"
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = ""
    protected = cell.paragraphs[0].add_run("FIELD_PLACEHOLDER")
    protected.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
    cell.paragraphs[0].add_run(" / ")
    out_of_scope = cell.paragraphs[0].add_run("OUT_OF_SCOPE")
    out_of_scope.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
    document.save(template)
    mapping = [{
        "field_id": "field",
        "label": "Field",
        "source_type": "fixed",
        "locators": [{
            "kind": "table_cell",
            "table_path": [{"table_index": 0, "row_index": 0, "column_index": 0}],
            "expected_text": "FIELD_PLACEHOLDER",
        }],
        "write_strategy": "replace_text",
        "output_format": "text",
    }]
    assert write_docx_by_mapping(template, output, mapping, {}) == []
    text = Document(output).tables[0].cell(0, 0).text
    assert "FIELD_PLACEHOLDER" in text
    assert "OUT_OF_SCOPE" not in text
    assert "Not assessed in MVP1" in text


def test_report_selector_keeps_periods_separate():
    docs = [
        AnnouncementDocument(
            announcement_id="annual",
            stock_code="000938",
            title="2025年年度报告",
            published_at=datetime(2026, 4, 1),
            url="https://example.com/annual.pdf",
            document_type="annual_report",
            report_year=2025,
        ),
        AnnouncementDocument(
            announcement_id="interim",
            stock_code="000938",
            title="2026年半年度报告",
            published_at=datetime(2026, 8, 30),
            url="https://example.com/interim.pdf",
            document_type="interim_report",
            report_year=2026,
        ),
    ]
    selected = select_reports(
        build_report_catalog(docs),
        as_of=date(2026, 9, 1),
        announcement_start=date(2026, 1, 1),
    )
    assert selected.latest_full_annual_report.announcement_id == "annual"
    assert selected.latest_interim_report.announcement_id == "interim"
    assert [item.announcement_id for item in selected.window_announcements] == ["interim", "annual"]


def test_cninfo_catalog_query_is_company_specific():
    def handler(request: httpx.Request) -> httpx.Response:
        assert b"stock=000938%2Cgssz0000938" in request.content
        return httpx.Response(
            200,
            json={
                "hasMore": False,
                "announcements": [{
                    "announcementId": "a1",
                    "announcementTitle": "2025年年度报告",
                    "announcementTime": 1775001600000,
                    "adjunctUrl": "finalpage/2026-04-01/report.pdf",
                    "secCode": "000938",
                }],
            },
        )

    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        documents = fetch_company_announcements(
            client,
            stock_code="000938",
            org_id="gssz0000938",
            start_date=date(2026, 1, 1),
            end_date=date(2026, 7, 3),
        )
    assert len(documents) == 1
    assert documents[0].document_type == "annual_report"
    assert documents[0].report_year == 2025
    assert documents[0].url.endswith("report.pdf")


def test_pdf_parser_reports_unapproved_dependency(monkeypatch, tmp_path):
    pdf = tmp_path / "report.pdf"
    pdf.write_bytes(b"%PDF-1.4\n")
    original_import = builtins.__import__

    def reject_pdfplumber(name, *args, **kwargs):
        if name == "pdfplumber":
            raise ImportError("not approved")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", reject_pdfplumber)
    with pytest.raises(DataSourceError, match="pdfplumber==0.11.10"):
        parse_machine_generated_pdf(pdf)


def test_pdf_parser_extracts_each_page_text_tables_and_page_numbers(tmp_path):
    assert version("pdfplumber") == "0.11.10"
    pdf = tmp_path / "machine-generated-report.pdf"
    pdf.write_bytes(_machine_generated_pdf())
    parsed = parse_machine_generated_pdf(pdf)
    assert [page.page_number for page in parsed.pages] == [1, 2]
    assert "Annual Report Page One" in parsed.pages[0].text
    assert "Annual Report Page Two" in parsed.pages[1].text
    assert parsed.pages[0].tables[0].rows == [["Metric", "Value"], ["Revenue", "100"]]
    assert parsed.pages[1].tables[0].rows == [["Item", "Amount"], ["Cash", "80"]]


def _self_check_state(tmp_path: Path, output: Path, mappings, failures):
    json_paths = {}
    for name in ("sources", "failed_fields", "extracted_data", "execution_plan"):
        path = tmp_path / f"{name}.json"
        path.write_text("[]", encoding="utf-8")
        json_paths[f"{name}_json_path"] = str(path)
    return {
        "output_docx_path": str(output),
        "field_mapping": mappings,
        "field_results": {},
        "evidence_records": [],
        "failed_fields": failures,
        "artifacts": [],
        "execution_plan": [],
        **json_paths,
    }


def test_self_check_verifies_failed_expected_empty_target(tmp_path):
    output = tmp_path / "result.docx"
    document = Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = ""
    document.save(output)
    mapping = [{
        "field_id": "empty_field",
        "label": "Empty Field",
        "source_type": "fixed",
        "locators": [{
            "kind": "table_cell",
            "table_path": [{"table_index": 0, "row_index": 0, "column_index": 0}],
            "expected_text": None,
            "expected_empty": True,
        }],
        "write_strategy": "replace_text",
        "output_format": "text",
    }]
    failures = [{"field_id": "empty_field", "reason": "No verified value"}]
    state = _self_check_state(tmp_path, output, mapping, failures)
    clean = self_check_node(state)
    assert clean["self_check_result"]["checks"]["failed_targets_preserved"] is True

    document = Document(output)
    document.tables[0].cell(0, 0).text = "unexpected"
    document.save(output)
    dirty = self_check_node(state)
    assert dirty["self_check_result"]["checks"]["failed_targets_preserved"] is False
    assert dirty["self_check_result"]["passed"] is False


def test_self_check_reports_corrupt_docx_without_raising(tmp_path):
    output = tmp_path / "result.docx"
    output.write_bytes(b"not a docx")
    state = _self_check_state(tmp_path, output, [], [])
    result = self_check_node(state)
    assert result["self_check_result"]["checks"]["docx_reopens"] is False
    assert result["self_check_result"]["passed"] is False
    assert any("DOCX cannot be reopened" in issue for issue in result["self_check_result"]["issues"])
