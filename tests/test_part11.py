from __future__ import annotations

from datetime import date, datetime, timezone
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from mlc_agent.audit_changes import (
    AuditRecord,
    EvidenceDocument,
    FactEvidence,
    MaterialChangeDisclosure,
    MaterialChangeEvent,
    Part11Input,
    RestatementDisclosure,
    collect_part11,
    collect_part11_groups,
    identify_big4,
)
from mlc_agent.config import load_template_mapping
from mlc_agent.docx_writer import validate_template_mapping
from mlc_agent.news import NewsArticle, select_and_deduplicate_news
from mlc_agent.production_adapters import (
    Part11BoardChangesExtraction,
    Part11RestatementExtraction,
)


AS_OF = date(2026, 7, 3)
CAPTURED = datetime(2026, 7, 3, 9, 0, tzinfo=timezone.utc)
ANNUAL = "https://static.cninfo.com.cn/annual-2025.pdf"
PRIOR = "https://static.cninfo.com.cn/annual-2024.pdf"
NOTICE = "https://static.cninfo.com.cn/notice.pdf"
EXCHANGE = "https://www.szse.cn/disclosure/penalty.html"
LATEST_TEXT = (
    "安永华明会计师事务所（特殊普通合伙）审计并出具标准无保留意见。"
    "本年度未发生审计机构变更。报告明确公司过去24个月不存在财务重述。"
    "公司明确披露董事及高管没有重大变化。公司业务经营没有重大变化。"
    "前三大股东没有重大变化。公司不存在应披露的重大未决诉讼。"
    "非四大会计师事务所背景为全国性证券审计机构。"
    "审计机构出具保留意见，保留事项涉及应收账款可收回性。"
    "2024年7月3日公司更正前期财务报表，影响2023年度收入确认。"
    "2026年2月1日公司董事长辞任，公司将其明确列为重大管理层变化。"
    "2020年案件仍为重大未决诉讼，涉案金额人民币5000万元。"
    "2023年案件已结案，曾为重大诉讼。"
)
PRIOR_TEXT = "安永华明会计师事务所（特殊普通合伙）审计2024年度财务报表。"
NOTICE_TEXT = "公司公告审计机构由甲所变更为乙所，原因为服务期限届满。"
EXCHANGE_TEXT = "交易所明确披露公司不存在监管调查或处罚。"


def fact(source, url, disclosure_date, period, text):
    return {
        "source": source, "source_url": url, "disclosure_date": disclosure_date,
        "period": period, "evidence_text": text, "page_number": 1,
    }


def matter_evidence(source, url, disclosure_date, text):
    return {
        "source": source, "source_url": url, "disclosure_date": disclosure_date,
        "evidence_text": text, "page_number": 1,
    }


@pytest.fixture
def documents():
    return [
        EvidenceDocument(source="annual_report", source_url=ANNUAL, disclosure_date=date(2026, 4, 20), period="FY2025", text=LATEST_TEXT),
        EvidenceDocument(source="annual_report", source_url=PRIOR, disclosure_date=date(2025, 4, 20), period="FY2024", text=PRIOR_TEXT),
        EvidenceDocument(source="cninfo", source_url=NOTICE, disclosure_date=date(2026, 5, 1), period="post-FY2025", text=NOTICE_TEXT),
        EvidenceDocument(source="exchange", source_url=EXCHANGE, disclosure_date=date(2026, 6, 1), period="through 2026-06-01", text=EXCHANGE_TEXT),
    ]


def news_articles():
    common = {
        "company_name": "测试公司", "published_date": "2026-03-01",
        "title_entities": ["测试公司", "监管处罚"], "fact_fingerprint": "因信息披露违规被处罚",
        "company_specific_adverse_fact": True,
    }
    return [
        NewsArticle(**common, publisher="原始媒体甲", language="zh", title="测试公司受到监管处罚",
                    original_url="https://media.example.com/a", source_kind="media",
                    body="监管机构因信息披露违规对测试公司作出处罚。",
                    english_summary="The regulator penalized 测试公司 for disclosure violations."),
        NewsArticle(**common, publisher="监管机构", language="en", title="测试公司处罚决定",
                    original_url="https://regulator.example.com/b", source_kind="regulator",
                    body=None, english_summary=None),
    ]


def base_input(documents):
    latest_evidence = fact("annual_report", ANNUAL, "2026-04-20", "FY2025", "安永华明会计师事务所（特殊普通合伙）审计并出具标准无保留意见。")
    prior_evidence = fact("annual_report", PRIOR, "2025-04-20", "FY2024", PRIOR_TEXT)
    no_board = fact("annual_report", ANNUAL, "2026-04-20", "FY2025", "公司明确披露董事及高管没有重大变化。")
    no_business = fact("annual_report", ANNUAL, "2026-04-20", "FY2025", "公司业务经营没有重大变化。")
    no_shareholder = fact("annual_report", ANNUAL, "2026-04-20", "FY2025", "前三大股东没有重大变化。")
    return Part11Input.model_validate({
        "documents": [item.model_dump(mode="json") for item in documents],
        "latest_audit": {
            "fiscal_year": 2025, "auditor_name": "安永华明会计师事务所（特殊普通合伙）",
            "opinion": "Unmodified", "evidence": latest_evidence,
        },
        "previous_audit": {
            "fiscal_year": 2024, "auditor_name": "安永华明会计师事务所（特殊普通合伙）",
            "opinion": "Unmodified", "evidence": prior_evidence,
        },
        "restatements": {
            "status": "no", "events": [],
            "negative_evidence": fact("annual_report", ANNUAL, "2026-04-20", "FY2025", "报告明确公司过去24个月不存在财务重述。"),
        },
        "board_officer_changes": {"status": "no", "events": [], "negative_evidence": no_board},
        "business_operation_changes": {"status": "no", "events": [], "negative_evidence": no_business},
        "top3_shareholder_changes": {"status": "no", "events": [], "negative_evidence": no_shareholder},
        "litigation": {
            "status": "yes",
            "matters": [{
                "kind": "litigation", "event_date": "2020-01-01", "subject": "测试公司",
                "amount": "5000", "currency": "CNY ten thousand", "status": "pending",
                "authority": "某市中级人民法院", "summary": "合同纠纷仍未决",
                "materiality_basis": "official_major", "pending_in_latest_report": True,
                "evidence": matter_evidence("annual_report", ANNUAL, "2026-04-20", "2020年案件仍为重大未决诉讼，涉案金额人民币5000万元。"),
            }, {
                "kind": "litigation", "event_date": "2023-01-01", "subject": "测试公司",
                "status": "closed", "authority": "某市法院", "summary": "历史案件已结案",
                "materiality_basis": "official_major", "pending_in_latest_report": False,
                "evidence": matter_evidence("annual_report", ANNUAL, "2026-04-20", "2023年案件已结案，曾为重大诉讼。"),
            }],
        },
        "regulatory": {
            "status": "no", "matters": [],
            "negative_evidence": matter_evidence("exchange", EXCHANGE, "2026-06-01", EXCHANGE_TEXT),
        },
        "news_articles": [item.model_dump(mode="json") for item in news_articles()],
    })


@pytest.mark.parametrize("name,expected", [
    ("德勤华永会计师事务所", "Deloitte"), ("PricewaterhouseCoopers Zhong Tian LLP", "PwC"),
    ("安永华明会计师事务所", "EY"), ("KPMG Huazhen LLP", "KPMG"), ("天健会计师事务所", None),
])
def test_big4_identification_is_limited_to_members(name, expected):
    assert identify_big4(name) == expected


def test_audit_profiles_and_modified_opinion_validation(documents, tmp_path):
    data = base_input(documents)
    result = collect_part11(data, company_name="测试公司", as_of=AS_OF, captured_at=CAPTURED, artifact_dir=tmp_path)
    values = {item.field_id: item.value for item in result.source_values}
    assert "Big 4: Yes (EY)" in values["auditor_profile"]
    assert values["audit_opinion"] == "Unmodified"
    assert values["qualified_opinion_details"] == "N/A - standard unqualified opinion"

    with pytest.raises(ValueError, match="requires details"):
        AuditRecord.model_validate({
            "fiscal_year": 2025, "auditor_name": "某会计师事务所", "opinion": "Qualified",
            "evidence": data.latest_audit.evidence.model_dump(mode="json"),
        })


def test_no_requires_explicit_negative_quotation(documents, tmp_path):
    data = base_input(documents)
    negative = data.restatements.negative_evidence
    assert negative is not None
    negative.evidence_text = "公司披露会计政策和会计估计。"
    result = collect_part11(
        data, company_name="测试公司", as_of=AS_OF, captured_at=CAPTURED, artifact_dir=tmp_path
    )
    assert {error.field_id for error in result.errors} == {
        "financial_restatement_status", "financial_restatement_details"
    }
    assert not any(
        value.field_id in {"financial_restatement_status", "financial_restatement_details"}
        for value in result.source_values
    )


def test_regulatory_no_requires_explicit_negative_quotation(documents):
    raw = base_input(documents).model_dump(mode="json")
    raw["regulatory"]["negative_evidence"]["evidence_text"] = "公司披露处罚及整改情况。"
    with pytest.raises(ValueError, match="explicitly negative quotation"):
        Part11Input.model_validate(raw)

def test_previous_audit_evidence_mismatch_does_not_discard_latest_audit_fields(
    documents, tmp_path
):
    data = base_input(documents)
    data.previous_audit.evidence.evidence_text = "text that is not in the supplied report"

    result = collect_part11(
        data,
        company_name="测试公司",
        as_of=AS_OF,
        captured_at=CAPTURED,
        artifact_dir=tmp_path,
    )

    values = {item.field_id: item.value for item in result.source_values}
    assert values["auditor_profile"].startswith("安永华明")
    assert values["audit_opinion"] == "Unmodified"
    assert "auditor_change_status" not in values
    assert any(item.field_id == "auditor_change_status" for item in result.errors)


def test_non_big4_requires_background_and_qualified_details(documents, tmp_path):
    data = base_input(documents)
    data.latest_audit = AuditRecord.model_validate({
        "fiscal_year": 2025, "auditor_name": "天健会计师事务所",
        "non_big4_background": "全国性证券审计机构", "opinion": "Qualified",
        "modified_opinion_details": "应收账款可收回性",
        "evidence": fact("annual_report", ANNUAL, "2026-04-20", "FY2025", "审计机构出具保留意见，保留事项涉及应收账款可收回性。"),
    })
    result = collect_part11(data, company_name="测试公司", as_of=AS_OF, captured_at=CAPTURED, artifact_dir=tmp_path)
    values = {item.field_id: item.value for item in result.source_values}
    assert "Big 4: No" in values["auditor_profile"]
    assert values["audit_opinion"] == "Qualified"
    assert values["qualified_opinion_details"] == "应收账款可收回性"


def test_restatement_24_month_boundary_is_inclusive(documents, tmp_path):
    data = base_input(documents)
    data.restatements = data.restatements.model_validate({
        "status": "yes", "events": [{
            "event_date": "2024-07-03", "periods_affected": "FY2023", "details": "收入确认更正",
            "evidence": fact("annual_report", ANNUAL, "2026-04-20", "FY2025", "2024年7月3日公司更正前期财务报表，影响2023年度收入确认。"),
        }]
    })
    result = collect_part11(data, company_name="测试公司", as_of=AS_OF, captured_at=CAPTURED, artifact_dir=tmp_path)
    values = {item.field_id: item.value for item in result.source_values}
    assert values["financial_restatement_status"] == "Yes"
    assert "2024-07-03" in values["financial_restatement_details"]


def test_old_pending_litigation_included_and_old_closed_filtered(documents, tmp_path):
    result = collect_part11(base_input(documents), company_name="测试公司", as_of=AS_OF, captured_at=CAPTURED, artifact_dir=tmp_path)
    details = next(item for item in result.source_values if item.field_id == "litigation_regulatory_details")
    assert "2020-01-01" in details.value
    assert "2023-01-01" not in details.value


def test_not_disclosed_is_verified_value_and_not_converted_to_no(documents, tmp_path):
    data = base_input(documents)
    data.litigation = data.litigation.model_validate({
        "status": "not_disclosed",
        "matters": [],
        "coverage_evidence": matter_evidence(
            "annual_report",
            ANNUAL,
            "2026-04-20",
            "2023年案件已结案，曾为重大诉讼。",
        ),
    })
    result = collect_part11(data, company_name="测试公司", as_of=AS_OF, captured_at=CAPTURED, artifact_dir=tmp_path)
    value = next(item for item in result.source_values if item.field_id == "litigation_status")
    assert value.value == "Not disclosed"
    assert not any(item.field_id == "litigation_status" for item in result.errors)


def test_material_change_requires_explicit_material_flag(documents):
    data = base_input(documents).model_dump(mode="json")
    data["board_officer_changes"] = {
        "status": "yes", "events": [{
            "category": "board_officer", "event_date": "2026-02-01", "details": "董事长辞任",
            "explicit_material": False,
            "evidence": fact("annual_report", ANNUAL, "2026-04-20", "FY2025", "2026年2月1日公司董事长辞任，公司将其明确列为重大管理层变化。"),
        }]
    }
    with pytest.raises(ValueError, match="explicit_material"):
        Part11Input.model_validate(data)


def test_news_deduplicates_event_but_preserves_support_sources_and_12m_window():
    articles = news_articles()
    articles.append(articles[0].model_copy(update={
        "published_date": date(2025, 7, 2), "original_url": "https://media.example.com/old"
    }))
    events = select_and_deduplicate_news(articles, as_of=AS_OF)
    assert len(events) == 1
    assert len(events[0].articles) == 2
    assert events[0].articles[0].body is not None


def test_search_engine_subdomain_cannot_be_news_evidence():
    raw = news_articles()[0].model_dump(mode="json")
    raw["original_url"] = "https://news.google.com/article/123"
    with pytest.raises(ValueError, match="discovery-only"):
        NewsArticle.model_validate(raw)


def test_additional_search_engine_domain_cannot_be_news_evidence():
    raw = news_articles()[0].model_dump(mode="json")
    raw["original_url"] = "https://news.sogou.com/article/123"
    with pytest.raises(ValueError, match="discovery-only"):
        NewsArticle.model_validate(raw)


def test_other_company_article_is_rejected_without_polluting_attachment(documents, tmp_path):
    data = base_input(documents)
    other = news_articles()[0].model_copy(update={
        "company_name": "另一家公司",
        "title": "另一家公司受到处罚",
        "title_entities": ["另一家公司", "监管处罚"],
        "original_url": "https://media.example.com/other",
        "body": "另一家公司因信息披露违规受到处罚。",
        "english_summary": "另一家公司 was penalized for disclosure violations.",
    })
    data.news_articles.append(other)
    result = collect_part11(
        data, company_name=" 测试公司 ", as_of=AS_OF, captured_at=CAPTURED, artifact_dir=tmp_path
    )
    mismatch = [item for item in result.errors if item.field_id == "negative_news_input"]
    assert len(mismatch) == 1 and "does not exactly match" in mismatch[0].reason
    summary = next(item.value for item in result.source_values if item.field_id == "negative_news_summary")
    assert "另一家公司" not in summary
    attachment_text = "\n".join(
        paragraph.text for paragraph in Document(result.artifacts[0].path).paragraphs
    )
    assert "另一家公司" not in attachment_text


def test_news_attachment_has_real_links_body_and_unavailable_rule(documents, tmp_path):
    result = collect_part11(base_input(documents), company_name="测试公司", as_of=AS_OF, captured_at=CAPTURED, artifact_dir=tmp_path)
    assert not [item for item in result.errors if item.field_id.startswith("negative_news")]
    path = Path(result.artifacts[0].path)
    assert path.name == "negative_news_articles.docx" and path.is_file()
    attachment = Document(path)
    text = "\n".join(item.text for item in attachment.paragraphs)
    assert "The regulator penalized 测试公司" in text
    assert "监管机构因信息披露违规" in text
    assert "Full text unavailable" in text
    assert len([rel for rel in attachment.part.rels.values() if rel.reltype == RT.HYPERLINK]) == 2
    with ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert "Rationale for Recommendation" not in document_xml
    assert "Premium" not in document_xml


def test_mapping_matches_template_and_contains_no_underwriting_fields():
    root = Path(__file__).resolve().parents[1]
    mapping = load_template_mapping(root / "configs")
    ids = {
        "auditor_profile", "audit_opinion", "qualified_opinion_details", "auditor_change_status",
        "auditor_change_details", "financial_restatement_status", "financial_restatement_details",
        "board_officer_material_change", "business_operation_material_change", "top3_shareholder_change",
        "nonfinancial_change_details", "litigation_status", "regulatory_status",
        "litigation_regulatory_details", "negative_news_summary",
    }
    fields = [item for item in mapping["fields"] if item["field_id"] in ids]
    assert len(fields) == 15
    assert not any(any(word in item["field_id"].lower() for word in ("recommend", "premium", "rationale")) for item in fields)
    validate_template_mapping(root / "Workup_template_260617-外测版.docx", fields)


def test_collect_part11_groups_accepts_plain_dict_documents():
    document = {
        "source": "annual_report",
        "source_url": ANNUAL,
        "disclosure_date": "2026-04-20",
        "period": "FY2025",
        "text": LATEST_TEXT,
        "pages": [],
    }
    extraction = Part11BoardChangesExtraction(
        board_officer_changes=MaterialChangeDisclosure(
            status="yes",
            events=[MaterialChangeEvent(
                category="board_officer",
                event_date=date(2026, 2, 1),
                details="公司董事长辞任，公司将其明确列为重大管理层变化。",
                explicit_material=True,
                evidence=FactEvidence(
                    source="annual_report", source_url=ANNUAL,
                    disclosure_date=date(2026, 4, 20), period="FY2025",
                    evidence_text="2026年2月1日公司董事长辞任，公司将其明确列为重大管理层变化。",
                    page_number=1,
                ),
            )],
        )
    )
    result = collect_part11_groups(
        documents=[document],
        matter_documents=[],
        extracted={"board_changes": extraction},
        groups={"board_changes"},
        company_name="紫光股份有限公司",
        as_of=date(2026, 6, 1),
        captured_at=CAPTURED,
        artifact_dir=Path("."),
    )
    assert any(value.field_id == "board_officer_material_change" for value in result.source_values)
    assert not any(error.field_id == "board_officer_material_change" for error in result.errors)


def test_restatement_weak_negative_quote_fails_only_its_two_fields():
    document = {
        "source": "annual_report", "source_url": ANNUAL,
        "disclosure_date": "2026-04-20", "period": "FY2025",
        "text": LATEST_TEXT, "pages": [],
    }
    weak = FactEvidence(
        source="annual_report", source_url=ANNUAL,
        disclosure_date=date(2026, 4, 20), period="FY2025",
        evidence_text="报告说明公司财务信息正常。",  # 不含 不存在/未发生/无重大 等明确否定词
        page_number=1,
    )
    extraction = Part11RestatementExtraction(
        restatements=RestatementDisclosure(status="no", negative_evidence=weak)
    )
    result = collect_part11_groups(
        documents=[document], matter_documents=[], extracted={"restatements": extraction},
        groups={"restatements"}, company_name="紫光股份有限公司",
        as_of=date(2026, 6, 1), captured_at=CAPTURED, artifact_dir=Path("."),
    )
    assert {error.field_id for error in result.errors} == {
        "financial_restatement_status", "financial_restatement_details"
    }
