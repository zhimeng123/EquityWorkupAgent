from pathlib import Path
import base64

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from mlc_agent.config import load_template_mapping
from mlc_agent.docx_writer import (
    NOT_ASSESSED_TEXT,
    find_unmarked_non_mvp_placeholders,
    validate_template_mapping,
    write_docx_by_mapping,
)
from mlc_agent.schemas import FieldResult


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "Workup_template_260617-外测版.docx"


def _fake_results(mappings, tmp_path):
    hyperlinks = {
        "official_website": "https://example.com",
        "xueqiu_url": "https://xueqiu.com/S/SZ000938",
        "eastmoney_url": "https://quote.eastmoney.com/sz000938.html",
    }
    image_path = tmp_path / "test.png"
    image_path.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    ))
    results = {
        item["field_id"]: FieldResult(
            field_id=item["field_id"],
            label=item["label"],
            value=hyperlinks.get(item["field_id"], f"TEST-{item['field_id']}"),
            selected_source="fixed",
            structured_value=(
                [[f"R{row}C{column}" for column in range(5)] for row in range(4)]
                if item["write_strategy"] == "fill_fixed_table"
                else None
            ),
            artifact_path=(str(image_path) if item["write_strategy"] == "insert_image" else None),
        )
        for item in mappings
    }
    return results


def test_fixed_template_mapping_matches_real_template():
    config = load_template_mapping(ROOT / "configs")
    assert len(config["fields"]) == len({item["field_id"] for item in config["fields"]})
    assert len(config["fields"]) > 16
    validate_template_mapping(TEMPLATE, config["fields"])


def test_write_all_mapped_positions_and_preserve_mvp1(tmp_path):
    mappings = load_template_mapping(ROOT / "configs")["fields"]
    output = tmp_path / "result.docx"
    results = _fake_results(mappings, tmp_path)
    failures = write_docx_by_mapping(TEMPLATE, output, mappings, results)
    assert failures == []
    assert output.exists()

    document = Document(output)
    assert "TEST-company_english_name" in document.sections[0].header.paragraphs[1].text
    assert "TEST-total_assets" in document.tables[1].cell(1, 1).text
    assert "TEST-week_52_high" in document.tables[10].cell(2, 3).text
    assert "TEST-business_description" in document.paragraphs[21].text
    assert "TEST-google_finance_url" in document.paragraphs[10].text
    assert "TEST-listed_subsidiaries" in document.tables[1].cell(5, 3).text
    assert "TEST-listed_outside_directorships" in document.tables[1].cell(6, 3).text
    assert find_unmarked_non_mvp_placeholders(document, mappings, results) == []

    hyperlink_targets = {
        relationship.target_ref
        for relationship in document.part.rels.values()
        if relationship.reltype.endswith("/hyperlink")
    }
    expected_hyperlinks = {
        result.value
        for field_id, result in results.items()
        if next(item for item in mappings if item["field_id"] == field_id)["write_strategy"]
        == "replace_with_hyperlink"
    }
    assert hyperlink_targets == expected_hyperlinks
    market_cap = results["market_capitalization"].value
    assert market_cap in document.tables[1].cell(4, 3).text
    assert market_cap in document.tables[10].cell(0, 3).text

    description_run = document.paragraphs[21].runs[0]
    assert description_run.font.all_caps is False
    assert description_run.font.name == "Arial"
    assert description_run._element.rPr.rFonts.get(qn("w:hAnsi")) == "Arial"
