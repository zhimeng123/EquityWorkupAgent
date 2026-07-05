from datetime import datetime, timezone
from pathlib import Path

from docx import Document
import httpx

from mlc_agent.config import load_yaml
from mlc_agent.docx_writer import write_docx_by_mapping
from mlc_agent.external_links import (
    build_google_finance_url,
    collect_external_links,
    discover_cninfo_org_id,
    google_finance_exchange,
    validate_cninfo_company_url,
)
from mlc_agent.schemas import CompanyIdentity, FieldResult


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "Workup_template_260617-外测版.docx"


def _company(suffix="SZ", exchange="深圳证券交易所", code="000938"):
    market = "0" if suffix in {"SZ", "BJ"} else "1"
    return CompanyIdentity(
        company_name="Test Company",
        company_short_name="Test",
        stock_code=code,
        exchange=exchange,
        eastmoney_secid=f"{market}.{code}",
        eastmoney_secu_code=f"{code}.{suffix}",
        xueqiu_symbol=f"{suffix}{code}",
    )


def test_google_finance_exchange_mapping_for_shenzhen_shanghai_and_beijing():
    assert google_finance_exchange(_company()) == "SHE"
    assert google_finance_exchange(_company("SH", "上海证券交易所", "600519")) == "SHA"
    assert google_finance_exchange(_company("BJ", "北京证券交易所", "430047")) is None
    assert build_google_finance_url(_company()) == (
        "https://www.google.com/finance/quote/000938:SHE?hl=en"
    )


def test_collect_validates_two_company_specific_direct_links():
    company = _company()

    def handler(request):
        if request.url.host == "www.google.com":
            assert request.url.path == "/finance/quote/000938:SHE"
            return httpx.Response(
                200,
                text="<html><title>Test Company (000938:SHE) - Google Finance</title></html>",
                request=request,
            )
        assert request.url.host == "www.cninfo.com.cn"
        assert request.url.params["stockCode"] == "000938"
        assert request.url.params["orgId"] == "gssz0000938"
        return httpx.Response(
            200,
            text="<html><h1>Test Company [000938]</h1></html>",
            request=request,
        )

    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        result = collect_external_links(
            client,
            company=company,
            cninfo_org_id="gssz0000938",
            captured_at=datetime(2026, 7, 3, tzinfo=timezone.utc),
        )
    assert result.errors == []
    by_id = {item.field_id: item for item in result.source_values}
    assert set(by_id) == {"google_finance_url", "cninfo_company_url"}
    assert "000938:SHE" in by_id["google_finance_url"].value
    assert "stockCode=000938" in by_id["cninfo_company_url"].value
    assert by_id["google_finance_url"].source_url == by_id["google_finance_url"].value
    assert by_id["cninfo_company_url"].source_url == by_id["cninfo_company_url"].value


def test_cninfo_wrong_company_response_is_rejected():
    def handler(request):
        return httpx.Response(
            200,
            text="<html><h1>Wrong Company [000001]</h1></html>",
            request=request,
        )

    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        try:
            validate_cninfo_company_url(client, _company(), org_id="gssz0000938")
        except ValueError as exc:
            assert "does not match stock code 000938" in str(exc)
        else:
            raise AssertionError("wrong-company CNINFO page was accepted")


def test_cninfo_org_id_discovery_accepts_only_one_exact_stock_code():
    def exact_handler(request):
        assert request.method == "POST"
        assert request.headers["content-type"].startswith("application/x-www-form-urlencoded")
        assert request.content == b"keyWord=000938&maxNum=10"
        assert not request.url.params
        return httpx.Response(
            200,
            json=[
                {"code": "000938", "orgId": "gssz0000938", "zwjc": "紫光股份"},
                {"code": "000001", "orgId": "gssz0000001", "zwjc": "平安银行"},
            ],
            request=request,
        )

    with httpx.Client(transport=httpx.MockTransport(exact_handler)) as client:
        assert discover_cninfo_org_id(client, _company()) == "gssz0000938"

    def ambiguous_handler(request):
        assert request.method == "POST"
        assert request.content == b"keyWord=000938&maxNum=10"
        return httpx.Response(
            200,
            json=[
                {"code": "000938", "orgId": "org-a"},
                {"code": "000938", "orgId": "org-b"},
            ],
            request=request,
        )

    with httpx.Client(transport=httpx.MockTransport(ambiguous_handler)) as client:
        try:
            discover_cninfo_org_id(client, _company())
        except ValueError as exc:
            assert "one exact match" in str(exc)
        else:
            raise AssertionError("ambiguous CNINFO orgId was accepted")


def test_beijing_google_link_failure_does_not_create_unverified_fallback():
    company = _company("BJ", "北京证券交易所", "430047")

    def handler(request):
        assert request.url.host == "www.cninfo.com.cn"
        return httpx.Response(200, text="Company [430047]", request=request)

    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        result = collect_external_links(
            client,
            company=company,
            cninfo_org_id="9900000001",
            captured_at=datetime(2026, 7, 3, tzinfo=timezone.utc),
        )
    assert [item.field_id for item in result.source_values] == ["cninfo_company_url"]
    assert result.errors == [
        "google_finance_url: Google Finance has no configured verifiable Beijing Stock Exchange page"
    ]


def test_part02_mapping_writes_two_clickable_external_relationships(tmp_path):
    mappings = load_yaml(ROOT / "configs" / "fields" / "part_02.yaml")["fields"]
    google = "https://www.google.com/finance/quote/000938:SHE?hl=en"
    cninfo = (
        "https://www.cninfo.com.cn/new/disclosure/stock?"
        "orgId=gssz0000938&stockCode=000938"
    )
    results = {
        "google_finance_url": FieldResult(
            field_id="google_finance_url",
            label="Google Finance Link",
            value=google,
            selected_source="system",
        ),
        "cninfo_company_url": FieldResult(
            field_id="cninfo_company_url",
            label="CNINFO Company Link",
            value=cninfo,
            selected_source="cninfo",
        ),
    }
    output = tmp_path / "links.docx"
    assert write_docx_by_mapping(TEMPLATE, output, mappings, results) == []
    document = Document(output)
    targets = {
        relationship.target_ref
        for relationship in document.part.rels.values()
        if relationship.reltype.endswith("/hyperlink")
    }
    assert targets == {google, cninfo}
    assert "000938" in document.paragraphs[10]._p.xml
    assert "000938" in document.paragraphs[12]._p.xml


def test_failed_links_leave_no_source_value_for_writer():
    def handler(request):
        return httpx.Response(200, text="wrong company 000001", request=request)

    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        result = collect_external_links(
            client,
            company=_company(),
            cninfo_org_id="wrong-org",
            captured_at=datetime(2026, 7, 3, tzinfo=timezone.utc),
        )
    assert result.source_values == []
    assert len(result.errors) == 2
