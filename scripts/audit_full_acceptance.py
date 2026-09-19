#!/usr/bin/env python3
"""Strict, external acceptance audit for one EquityWorkupAgent run.

The production workflow remains unchanged.  This script audits only persisted
artifacts and the final DOCX, so a successful CLI exit cannot hide a partial
run.  Rendering is structural; visual inspection is deliberately represented
by a separate, human-edited manifest.
"""

from __future__ import annotations

import argparse
import json
import subprocess
import sys
from pathlib import Path
from typing import Any

from docx import Document

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from mlc_agent.config import load_template_mapping
from mlc_agent.nodes import _fixed_table_is_written, _value_is_written


EXPECTED_FIELD_COUNT = 95
REQUIRED_ARTIFACTS = (
    "sources.json",
    "extracted_data.json",
    "failed_fields.json",
    "execution_plan.json",
    "result.docx",
)
RENDERER = (
    Path.home()
    / ".codex/plugins/cache/openai-primary-runtime/documents/26.909.12148/skills/documents/render_docx.py"
)


def _load_json(path: Path) -> Any:
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def _check(condition: bool, name: Any = None, details: Any = None) -> dict[str, Any]:
    result: dict[str, Any] = {"passed": bool(condition)}
    if details is not None:
        result["details"] = details
    elif name is not None:
        result["details"] = name
    return result


def _docx_field_check(document: Any, mapping: dict[str, Any], evidence: dict[str, Any]) -> bool:
    value = evidence.get("normalized_value", evidence.get("value"))
    if value is None:
        return False
    value_text = str(value)
    if mapping["write_strategy"] == "fill_fixed_table":
        # sources.json stores the rendered table as text.  Validate every
        # non-empty serialized row token inside the mapped table cells.
        for locator in mapping["locators"]:
            cell = None
            table = None
            for depth, step in enumerate(locator["table_path"]):
                tables = document.tables if depth == 0 else cell.tables
                table = tables[step["table_index"]]
                cell = table.cell(step["row_index"], step["column_index"])
            table_text = "\n".join(cell.text for row in table.rows for cell in row.cells)
            tokens = [
                token.strip()
                for line in value_text.splitlines()
                if line.strip()
                for token in line.split("|")
                if token.strip()
            ]
            if not tokens or any(token not in table_text for token in tokens):
                return False
        return True
    return _value_is_written(document, mapping, value_text)


def _render(docx_path: Path, render_dir: Path) -> dict[str, Any]:
    render_dir.mkdir(parents=True, exist_ok=True)
    bundled_python = (
        Path.home()
        / ".cache/codex-runtimes/codex-primary-runtime/dependencies/python/bin/python3"
    )
    renderer_python = bundled_python if bundled_python.is_file() else Path(sys.executable)
    command = [str(renderer_python), str(RENDERER), str(docx_path), "--output_dir", str(render_dir)]
    completed = subprocess.run(command, capture_output=True, text=True, check=False)
    pages = sorted(render_dir.glob("page-*.png"))
    result = {
        "command": command,
        "returncode": completed.returncode,
        "stdout": completed.stdout[-4000:],
        "stderr": completed.stderr[-4000:],
        "pages": [
            {"page": int(path.stem.split("-")[-1]), "path": str(path.resolve()), "bytes": path.stat().st_size}
            for path in pages
        ],
    }
    (render_dir / "render_manifest.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    template = {
        "pages": [
            {
                "page": item["page"],
                "status": "pending",
                "notes": "Inspect clipping, overlap, tables, fonts, headers/footers, and page breaks.",
            }
            for item in result["pages"]
        ]
    }
    (render_dir / "visual_inspection.template.json").write_text(
        json.dumps(template, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return result


def _visual_check(render_result: dict[str, Any], manifest_path: Path | None) -> dict[str, Any]:
    if manifest_path is None:
        return _check(False, "visual_inspection_manifest_present", "required for strict acceptance")
    try:
        manifest = _load_json(manifest_path)
    except (OSError, json.JSONDecodeError) as exc:
        return _check(False, "visual_inspection_manifest_valid", str(exc))
    pages = {item["page"] for item in render_result["pages"]}
    entries = manifest.get("pages") if isinstance(manifest, dict) else None
    if not isinstance(entries, list):
        return _check(False, "visual_inspection_manifest_valid", "pages must be a list")
    inspected = {
        item.get("page")
        for item in entries
        if isinstance(item, dict) and item.get("status") == "pass" and str(item.get("notes", "")).strip()
    }
    return _check(
        pages and inspected == pages and len(entries) == len(pages),
        "all_rendered_pages_visually_inspected",
        {"rendered_pages": sorted(pages), "passed_pages": sorted(inspected)},
    )


def audit_run(run_dir: Path, *, render: bool, inspection_manifest: Path | None) -> dict[str, Any]:
    run_dir = run_dir.resolve()
    checks: dict[str, dict[str, Any]] = {}
    issues: list[str] = []
    paths = {name: run_dir / name for name in REQUIRED_ARTIFACTS}
    checks["required_artifacts_present"] = _check(
        all(path.is_file() and path.stat().st_size > 0 for path in paths.values()),
        "required files are present and non-empty",
    )

    config = load_template_mapping(run_dir.parent.parent / "configs")
    mappings = config["fields"]
    configured_ids = {item["field_id"] for item in mappings}
    checks["configured_field_count_is_95"] = _check(
        len(mappings) == EXPECTED_FIELD_COUNT and len(configured_ids) == EXPECTED_FIELD_COUNT,
        {"count": len(mappings), "unique_count": len(configured_ids)},
    )

    loaded: dict[str, Any] = {}
    for name, path in paths.items():
        if path.suffix == ".json" and path.exists():
            try:
                loaded[name] = _load_json(path)
            except (OSError, json.JSONDecodeError) as exc:
                checks[f"{name}_parses"] = _check(False, str(exc))
    sources = loaded.get("sources.json")
    checks["sources_json_is_object"] = _check(
        isinstance(sources, dict),
        "sources.json must contain an object",
    )
    evidence = sources.get("evidence") if isinstance(sources, dict) else None
    checks["sources_evidence_is_list"] = _check(
        isinstance(evidence, list),
        "sources.json evidence must be a list",
    )
    evidence = evidence if isinstance(evidence, list) else []
    evidence_ids = [
        item.get("field_id")
        for item in evidence
        if isinstance(item, dict) and item.get("field_id")
    ]
    duplicate_evidence_ids = sorted(
        field_id
        for field_id in set(evidence_ids)
        if evidence_ids.count(field_id) > 1
    )
    evidence_by_id = {
        item.get("field_id"): item
        for item in evidence
        if isinstance(item, dict) and item.get("field_id")
    }
    checks["evidence_field_ids_are_unique_and_nonempty"] = _check(
        len(evidence_ids) == len(evidence)
        and not duplicate_evidence_ids,
        {"duplicate_ids": duplicate_evidence_ids},
    )
    failed = loaded.get("failed_fields.json")
    checks["failed_fields_json_is_list"] = _check(
        isinstance(failed, list),
        "failed_fields.json must contain a list",
    )
    failed = failed if isinstance(failed, list) else []
    failed_ids = {item.get("field_id") for item in failed if isinstance(item, dict)}
    checks["success_count_is_95"] = _check(
        len(evidence) == EXPECTED_FIELD_COUNT and set(evidence_by_id) == configured_ids,
        {"success_count": len(evidence), "missing": sorted(configured_ids - set(evidence_by_id)), "extra": sorted(set(evidence_by_id) - configured_ids)},
    )
    checks["failed_count_is_0"] = _check(
        isinstance(loaded.get("failed_fields.json"), list) and not failed,
        {"failed_count": len(failed), "failed_ids": sorted(failed_ids)},
    )
    conflicts = sources.get("conflicts") if isinstance(sources, dict) else None
    checks["conflicts_is_list"] = _check(
        isinstance(conflicts, list),
        "sources.json conflicts must be a list",
    )
    checks["conflict_count_is_0"] = _check(
        isinstance(conflicts, list) and not conflicts,
        {"conflict_count": len(conflicts) if isinstance(conflicts, list) else None},
    )
    checks["every_success_has_nonempty_value"] = _check(
        all(
            isinstance(item, dict)
            and str(item.get("normalized_value", item.get("value", ""))).strip()
            for item in evidence
        ),
        [item.get("field_id") for item in evidence if not isinstance(item, dict) or not str(item.get("normalized_value", item.get("value", ""))).strip()],
    )
    checks["every_success_has_source_url"] = _check(
        all(isinstance(item, dict) and str(item.get("source_url", "")).strip() for item in evidence),
        [item.get("field_id") for item in evidence if not isinstance(item, dict) or not str(item.get("source_url", "")).strip()],
    )

    extracted = loaded.get("extracted_data.json", {})
    checks["extracted_data_has_run_and_company"] = _check(
        isinstance(extracted, dict) and isinstance(extracted.get("run"), dict) and isinstance(extracted.get("company"), dict),
    )
    node_errors = extracted.get("node_errors") if isinstance(extracted, dict) else None
    checks["extracted_data_node_errors_is_list"] = _check(
        isinstance(node_errors, list),
        "extracted_data.node_errors must be a list",
    )
    checks["extracted_data_has_no_node_errors"] = _check(
        isinstance(node_errors, list) and not node_errors,
        {"node_error_count": len(node_errors) if isinstance(node_errors, list) else None},
    )
    execution_plan = loaded.get("execution_plan.json", [])
    checks["execution_plan_is_list"] = _check(
        isinstance(execution_plan, list),
        "execution_plan.json must contain a list",
    )
    execution_plan = execution_plan if isinstance(execution_plan, list) else []
    terminal_steps = {
        item.get("step_id"): item.get("status")
        for item in execution_plan
        if isinstance(item, dict) and item.get("step_id") in {"self_check", "finalize"}
    }
    checks["self_check_and_finalize_completed"] = _check(
        terminal_steps == {"self_check": "completed", "finalize": "completed"},
        terminal_steps,
    )

    docx_ok = False
    docx_checks: list[str] = []
    try:
        document = Document(paths["result.docx"])
        docx_ok = True
        mapping_by_id = {item["field_id"]: item for item in mappings}
        bad_fields = [
            field_id
            for field_id, item in evidence_by_id.items()
            if field_id not in mapping_by_id or not _docx_field_check(document, mapping_by_id[field_id], item)
        ]
        checks["all_95_values_match_docx_mapping"] = _check(
            len(evidence) == EXPECTED_FIELD_COUNT
            and set(evidence_by_id) == configured_ids
            and not bad_fields,
            {
                "mismatched_fields": bad_fields,
                "success_count": len(evidence),
                "missing": sorted(configured_ids - set(evidence_by_id)),
                "extra": sorted(set(evidence_by_id) - configured_ids),
            },
        )
    except Exception as exc:
        docx_checks.append(str(exc))
        checks["all_95_values_match_docx_mapping"] = _check(False, docx_checks)
    checks["docx_reopens"] = _check(docx_ok)

    render_result = {"pages": []}
    if render:
        render_result = _render(paths["result.docx"], run_dir / "rendered")
        checks["render_completed"] = _check(
            render_result["returncode"] == 0 and bool(render_result["pages"]) and all(item["bytes"] > 0 for item in render_result["pages"]),
            {"page_count": len(render_result["pages"]), "returncode": render_result["returncode"]},
        )
        checks["visual_inspection"] = _visual_check(render_result, inspection_manifest)
    else:
        checks["render_completed"] = _check(False, "rerun with --render")
        checks["visual_inspection"] = _check(False, "rerun with --render and --inspection-manifest")

    for name, result in checks.items():
        if not result["passed"]:
            issues.append(name)
    report = {
        "status": "PASS" if not issues else "FAIL",
        "run_dir": str(run_dir),
        "checks": checks,
        "issues": issues,
        "render": render_result,
        "evidence_contract": {
            "configured_fields": EXPECTED_FIELD_COUNT,
            "successes": len(evidence),
            "failures": len(failed),
            "conflicts": len(conflicts) if isinstance(conflicts, list) else None,
        },
    }
    (run_dir / "acceptance_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description="Strict external 95-field acceptance audit")
    parser.add_argument("run_dir", type=Path)
    parser.add_argument("--render", action="store_true", help="render result.docx to run_dir/rendered")
    parser.add_argument("--inspection-manifest", type=Path, help="JSON completed after inspecting every rendered page")
    args = parser.parse_args()
    report = audit_run(args.run_dir, render=args.render, inspection_manifest=args.inspection_manifest)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
