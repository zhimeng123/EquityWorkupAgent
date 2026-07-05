from __future__ import annotations

import argparse
from collections import defaultdict
from decimal import Decimal
import json
from pathlib import Path
import shutil

from docx import Document
from docx.shared import Pt, RGBColor

from mlc_agent.audit_changes import identify_big4
from mlc_agent.config import load_template_mapping
from mlc_agent.docx_writer import write_docx_by_mapping
from mlc_agent.litigation import LegalRegulatoryMatter, format_matters
from mlc_agent.operating_performance import BreakdownItem, _format_breakdown
from mlc_agent.schemas import FieldResult


def _related_party_summary(source_value: dict) -> list[str]:
    items = list(source_value.get("raw_value") or [])
    totals: dict[str, Decimal] = defaultdict(Decimal)
    counts: dict[str, int] = defaultdict(int)
    for item in items:
        kind = str(item["transaction_type"])
        counts[kind] += 1
        if item.get("amount_cny") is not None:
            totals[kind] += Decimal(str(item["amount_cny"]))
    lines = [
        f"{len(items)} verified related-party transaction records were extracted from FY2025 annual-report disclosures.",
        "Aggregate by transaction type:",
    ]
    for kind in sorted(counts, key=lambda key: (totals[key], counts[key]), reverse=True):
        lines.append(
            f"- {kind}: {counts[kind]} records; CNY {totals[kind] / Decimal('1000000'):,.2f} million."
        )
    ranked = sorted(
        (item for item in items if item.get("amount_cny") is not None),
        key=lambda item: Decimal(str(item["amount_cny"])),
        reverse=True,
    )[:15]
    lines.append("Largest 15 disclosed records:")
    for index, item in enumerate(ranked, start=1):
        amount = Decimal(str(item["amount_cny"])) / Decimal("1000000")
        lines.append(
            f"{index}. {item['related_party']} ({item['relationship']}) - "
            f"{item['transaction_type']}; CNY {amount:,.2f} million; "
            f"Terms: {item['commercial_terms']}; Board approval: {item['board_approved']}."
        )
    lines.append("The complete 147-item schedule remains in the supporting validated JSON evidence.")
    return lines


def _split_related_party_paragraphs(
    docx_path: Path, mappings: list[dict], lines: list[str]
) -> None:
    """Avoid renderer corruption from one 16k-character paragraph in a table cell."""
    mapping = next(
        item for item in mappings if item["field_id"] == "related_party_transactions"
    )
    locator = mapping["locators"][0]
    document = Document(docx_path)
    cell = None
    for depth, step in enumerate(locator["table_path"]):
        tables = document.tables if depth == 0 else cell.tables
        cell = tables[step["table_index"]].cell(step["row_index"], step["column_index"])
    assert cell is not None
    cell.text = ""
    for index, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1
        run = paragraph.add_run(line)
        run.font.name = "Arial"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
    document.save(docx_path)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Write validated local annual-report fields into an existing result DOCX"
    )
    parser.add_argument("result_docx", type=Path)
    parser.add_argument("annual_fields_json", type=Path)
    parser.add_argument("related_parties_json", type=Path)
    parser.add_argument("--config-dir", type=Path, default=Path("configs"))
    args = parser.parse_args()

    annual = json.loads(args.annual_fields_json.read_text(encoding="utf-8"))
    related = json.loads(args.related_parties_json.read_text(encoding="utf-8"))
    mapping_config = load_template_mapping(args.config_dir)
    mappings = mapping_config["fields"]
    labels = {item["field_id"]: item["label"] for item in mappings}
    results: dict[str, FieldResult] = {}

    def add(field_id: str, value: str, *, source: str = "annual_report", structured=None) -> None:
        results[field_id] = FieldResult(
            field_id=field_id,
            label=labels[field_id],
            value=value,
            selected_source=source,
            structured_value=structured,
        )

    for item in annual["source_values"]:
        field_id = item["field_id"]
        if field_id == "revenue_breakdown":
            breakdown = [BreakdownItem.model_validate(value) for value in item["value"]]
            add(field_id, _format_breakdown(breakdown), structured=item["value"])
        else:
            add(field_id, str(item["value"]), source=item.get("source", "annual_report"))

    related_value = related.get("source_value")
    if related_value:
        add("related_party_transactions", "\n".join(_related_party_summary(related_value)))

    groups = annual["part11_extractions"]
    audit = groups["audit"]["latest_audit"]
    big4 = identify_big4(audit["auditor_name"])
    auditor_profile = (
        f'{audit["auditor_name"]} | Big 4: Yes ({big4})'
        if big4
        else audit["auditor_name"]
    )
    add("auditor_profile", auditor_profile)
    add("audit_opinion", audit["opinion"])
    add(
        "qualified_opinion_details",
        audit.get("modified_opinion_details") or "N/A - standard unqualified opinion",
    )

    restatement = groups["restatements"]["restatements"]
    add("financial_restatement_status", "No" if restatement["status"] == "no" else "Yes")
    add(
        "financial_restatement_details",
        "Not applicable" if not restatement.get("events") else "\n".join(
            event["details"] for event in restatement["events"]
        ),
    )

    board = groups["board_changes"]["board_officer_changes"]
    board_events = board.get("events", [])
    add("board_officer_material_change", "Yes" if board_events else "No")
    if board_events:
        add(
            "nonfinancial_change_details",
            "\n".join(
                f'{event["event_date"]} | board_officer | {event["details"]}'
                for event in board_events
            ),
            source="cninfo",
        )

    litigation = groups["litigation"]["litigation"]
    regulatory = groups["regulatory"]["regulatory"]
    litigation_matters = [
        LegalRegulatoryMatter.model_validate(item) for item in litigation.get("matters", [])
    ]
    regulatory_matters = [
        LegalRegulatoryMatter.model_validate(item) for item in regulatory.get("matters", [])
    ]
    add("litigation_status", "Yes" if litigation_matters else "No")
    add("regulatory_status", "Yes" if regulatory_matters else "No")
    all_matters = litigation_matters + regulatory_matters
    add(
        "litigation_regulatory_details",
        format_matters(all_matters) if all_matters else "Not applicable",
        source="cninfo",
    )

    backup = args.result_docx.with_name("result_before_annual_writeback.docx")
    shutil.copy2(args.result_docx, backup)
    staged = args.result_docx.with_name("result_annual_writeback_staged.docx")
    failures = write_docx_by_mapping(args.result_docx, staged, mappings, results)
    if failures:
        raise RuntimeError(f"DOCX write failures: {failures}")
    if related_value:
        _split_related_party_paragraphs(
            staged, mappings, _related_party_summary(related_value)
        )
    staged.replace(args.result_docx)
    print(json.dumps({
        "result_docx": str(args.result_docx.resolve()),
        "backup_docx": str(backup.resolve()),
        "written_fields": sorted(results),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
